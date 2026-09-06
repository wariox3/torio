from datetime import date
from decimal import Decimal
from io import BytesIO

from django_tenants.test.cases import TenantTestCase
from openpyxl import load_workbook
from rest_framework import permissions
from rest_framework.test import APIRequestFactory

from general.models import (
    GenDocumento,
    GenDocumentoDetalle,
    GenDocumentoTipo,
    GenItem,
)
from inventario.models import InvAlmacen, InvExistencia
from inventario.views.informe import InvInformeViewSet


class _InformeViewSinPermisos(InvInformeViewSet):
    """Variante de la vista sin auth/permiso/throttle para probar los actions aislados."""
    authentication_classes = []
    permission_classes = [permissions.AllowAny]
    throttle_classes = []


class _InformeBase(TenantTestCase):
    """
    Escenario compartido por los cuatro informes: el mismo inventario visto como
    saldo consolidado, como saldo por almacén, valorizado y como movimiento.

        Martillo   inventario, 10 en Principal (4 remisionadas) + 5 en Bodega
        Clavo      inventario, saldo en cero
        Tornillo   inventario, -3 en Bodega
        Asesoría   servicio, sin inventario

    Los saldos se escriben a mano: acá se prueba el informe, no el servicio que
    los mueve. Por eso el consolidado del item se fija igual a la suma de sus
    filas de almacén, y las líneas de documento igual al consolidado, que son las
    dos invariantes que `_afectar_inventario` mantiene.
    """

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '+573000000000'
        tenant.correo = 'test@test.com'

    def setUp(self):
        self.principal = InvAlmacen.objects.create(nombre='Principal')
        self.bodega = InvAlmacen.objects.create(nombre='Bodega')

        self.martillo = self._item(
            'Martillo', 'M1', costo='900', costo_promedio='1000', precio='1500',
            existencia='15', remision='4', disponible='11',
        )
        self.clavo = self._item('Clavo', 'C1', costo_promedio='2')
        self.tornillo = self._item('Tornillo', 'T1', existencia='-3', disponible='-3')
        self.asesoria = self._item('Asesoría', 'A1', inventario=False)

        self.martillo_principal = self._saldo(
            self.martillo, self.principal, existencia='10', remision='4', disponible='6',
        )
        self.martillo_bodega = self._saldo(
            self.martillo, self.bodega, existencia='5', disponible='5',
        )
        self.clavo_principal = self._saldo(self.clavo, self.principal)
        self.tornillo_bodega = self._saldo(
            self.tornillo, self.bodega, existencia='-3', disponible='-3',
        )

        self.compra = self._documento_tipo('Compra', operacion_inventario=1)
        self.factura = self._documento_tipo('Factura', operacion_inventario=-1)
        self.remision_tipo = self._documento_tipo('Remisión', operacion_remision=-1)
        self.recibo = self._documento_tipo('Recibo de caja')

        # 18 martillos entran, 3 salen: los 15 del consolidado.
        self.entrada = self._movimiento(
            self.compra, date(2026, 1, 10), self.martillo, self.principal,
            cantidad='12', cantidad_operada='12', operacion_inventario=1,
        )
        self.entrada_bodega = self._movimiento(
            self.compra, date(2026, 1, 11), self.martillo, self.bodega,
            cantidad='6', cantidad_operada='6', operacion_inventario=1,
        )
        self.salida = self._movimiento(
            self.factura, date(2026, 1, 20), self.martillo, self.principal,
            cantidad='3', cantidad_operada='-3', operacion_inventario=-1,
        )
        # Despacho: mueve remisión, no existencia.
        self.despacho = self._movimiento(
            self.remision_tipo, date(2026, 1, 21), self.martillo, self.principal,
            cantidad='4', cantidad_operada='-4', operacion_remision=-1,
        )

        # Tres líneas que NO movieron inventario, una por cada mitad de la
        # invariante: sin aprobar, sin almacén y sin operación.
        self.sin_aprobar = self._movimiento(
            self.compra, date(2026, 1, 22), self.martillo, self.principal,
            cantidad='99', cantidad_operada='99', operacion_inventario=1,
            aprobado=False,
        )
        self.sin_almacen = self._movimiento(
            self.compra, date(2026, 1, 23), self.martillo, None,
            cantidad='7', cantidad_operada='7', operacion_inventario=1,
        )
        self.sin_operacion = self._movimiento(
            self.recibo, date(2026, 1, 24), self.martillo, self.principal,
            cantidad='5', cantidad_operada='5',
        )

    @staticmethod
    def _item(nombre, codigo, costo='0', costo_promedio='0', precio='0', inventario=True,
              existencia='0', remision='0', disponible='0'):
        return GenItem.objects.create(
            nombre=nombre, codigo=codigo, referencia=f'REF-{codigo}',
            producto=inventario, servicio=not inventario, inventario=inventario,
            costo=Decimal(costo),
            costo_promedio=Decimal(costo_promedio),
            costo_total=Decimal(costo_promedio) * Decimal(existencia),
            precio=Decimal(precio),
            existencia=Decimal(existencia),
            remision=Decimal(remision),
            disponible=Decimal(disponible),
        )

    @staticmethod
    def _saldo(item, almacen, existencia='0', remision='0', disponible='0'):
        return InvExistencia.objects.create(
            item=item, almacen=almacen,
            existencia=Decimal(existencia),
            remision=Decimal(remision),
            disponible=Decimal(disponible),
        )

    @staticmethod
    def _documento_tipo(nombre, operacion_inventario=0, operacion_remision=0):
        return GenDocumentoTipo.objects.create(
            nombre=nombre,
            inventario=bool(operacion_inventario or operacion_remision),
            operacion_inventario=operacion_inventario,
            operacion_remision=operacion_remision,
        )

    def _movimiento(self, documento_tipo, fecha, item, almacen, cantidad='0',
                    cantidad_operada='0', operacion_inventario=0, operacion_remision=0,
                    aprobado=True):
        documento = GenDocumento.objects.create(
            documento_tipo=documento_tipo,
            numero=GenDocumento.objects.count() + 1,
            fecha=fecha,
            estado_aprobado=aprobado,
        )
        return GenDocumentoDetalle.objects.create(
            documento=documento,
            item=item,
            almacen=almacen,
            cantidad=Decimal(cantidad),
            cantidad_operada=Decimal(cantidad_operada),
            operacion_inventario=operacion_inventario,
            operacion_remision=operacion_remision,
        )

    def _post(self, accion, **payload):
        payload.setdefault('informe', self.informe)
        request = APIRequestFactory().post(
            f'/inventario/informe/{accion}/', payload, format='json',
        )
        return _InformeViewSinPermisos.as_view({'post': accion})(request)

    def _lista(self, **payload):
        response = self._post('lista', **payload)
        self.assertEqual(response.status_code, 200)
        return response.data['results']

    def _ids(self, **payload):
        return {fila['id'] for fila in self._lista(**payload)}


class ExistenciaTests(_InformeBase):
    """Una fila por item que maneja inventario, con el saldo de todos los almacenes."""

    informe = 'existencia'

    def test_trae_los_items_de_inventario_incluido_el_que_esta_en_cero(self):
        """
        Un item que nunca movió sale en cero: el informe muestra el catálogo con
        su saldo, no solo lo que hay en bodega. Para ver lo que tiene saldo está
        `filtros`.
        """
        self.assertEqual(
            self._ids(), {self.martillo.id, self.clavo.id, self.tornillo.id},
        )

    def test_deja_por_fuera_lo_que_no_maneja_inventario(self):
        self.assertNotIn(self.asesoria.id, self._ids())

    def test_trae_los_tres_saldos(self):
        fila = next(f for f in self._lista() if f['id'] == self.martillo.id)

        self.assertEqual(fila['codigo'], 'M1')
        self.assertEqual(fila['nombre'], 'Martillo')
        self.assertEqual(fila['referencia'], 'REF-M1')
        self.assertEqual(Decimal(fila['existencia']), Decimal('15'))
        self.assertEqual(Decimal(fila['remision']), Decimal('4'))
        self.assertEqual(Decimal(fila['disponible']), Decimal('11'))

    def test_no_trae_costo(self):
        """El costo es lo único que separa este informe de `inventario_valorizado`."""
        fila = next(f for f in self._lista() if f['id'] == self.martillo.id)

        self.assertNotIn('costo_promedio', fila)
        self.assertNotIn('costo_total', fila)

    def test_el_negativo_no_se_oculta(self):
        self.assertIn(self.tornillo.id, self._ids())

    def test_ordena_por_nombre_por_defecto(self):
        self.assertEqual(
            [fila['nombre'] for fila in self._lista()],
            ['Clavo', 'Martillo', 'Tornillo'],
        )

    def test_el_filtro_del_cliente_no_puede_ampliar_la_invariante(self):
        """
        Los `filtros` se aplican sobre el queryset ya filtrado por el informe, así
        que se intersecan con la invariante: pedir `inventario = false` no trae la
        asesoría, no trae nada.
        """
        self.assertEqual(
            self._lista(
                filtros=[{'propiedad': 'inventario', 'operador': '=', 'valor': False}],
            ),
            [],
        )

    def test_filtra_por_saldo(self):
        self.assertEqual(
            self._ids(filtros=[{'propiedad': 'existencia', 'operador': '!=', 'valor': 0}]),
            {self.martillo.id, self.tornillo.id},
        )


class ExistenciaAlmacenTests(_InformeBase):
    """
    El mismo informe abierto por almacén: la fila es el saldo de un item en un
    almacén, que es lo que expone `/inventario/existencia/` en itrio.
    """

    informe = 'existencia_almacen'

    def test_trae_una_fila_por_item_y_almacen(self):
        self.assertEqual(
            self._ids(),
            {
                self.martillo_principal.id,
                self.martillo_bodega.id,
                self.clavo_principal.id,
                self.tornillo_bodega.id,
            },
        )

    def test_el_desglose_suma_el_consolidado_del_item(self):
        """
        Es la invariante que mantiene `_afectar_inventario`, y la razón por la que
        los dos informes son el mismo informe a dos granularidades.
        """
        filas = [f for f in self._lista() if f['item_id'] == self.martillo.id]

        self.assertEqual(len(filas), 2)
        self.assertEqual(
            sum(Decimal(f['existencia']) for f in filas), self.martillo.existencia,
        )
        self.assertEqual(
            sum(Decimal(f['disponible']) for f in filas), self.martillo.disponible,
        )

    def test_las_columnas_del_item_y_del_almacen_vienen_resueltas(self):
        fila = next(f for f in self._lista() if f['id'] == self.martillo_principal.id)

        self.assertEqual(fila['item_codigo'], 'M1')
        self.assertEqual(fila['item_nombre'], 'Martillo')
        self.assertEqual(fila['item_referencia'], 'REF-M1')
        self.assertEqual(fila['almacen_id'], self.principal.id)
        self.assertEqual(fila['almacen_nombre'], 'Principal')
        self.assertEqual(Decimal(fila['existencia']), Decimal('10'))
        self.assertEqual(Decimal(fila['remision']), Decimal('4'))
        self.assertEqual(Decimal(fila['disponible']), Decimal('6'))
        self.assertEqual(Decimal(fila['costo_promedio']), Decimal('1000'))

    def test_ordena_por_almacen_y_despues_por_item(self):
        self.assertEqual(
            [(f['almacen_nombre'], f['item_nombre']) for f in self._lista()],
            [
                ('Bodega', 'Martillo'),
                ('Bodega', 'Tornillo'),
                ('Principal', 'Clavo'),
                ('Principal', 'Martillo'),
            ],
        )

    def test_filtra_por_almacen_y_cruzando_el_item(self):
        self.assertEqual(
            self._ids(filtros=[
                {'propiedad': 'almacen_id', 'operador': '=', 'valor': self.bodega.id},
            ]),
            {self.martillo_bodega.id, self.tornillo_bodega.id},
        )
        self.assertEqual(
            self._ids(filtros=[
                {'propiedad': 'item__nombre', 'operador': 'contiene', 'valor': 'marti'},
            ]),
            {self.martillo_principal.id, self.martillo_bodega.id},
        )


class InventarioValorizadoTests(_InformeBase):
    """`existencia` con lo que vale cada saldo."""

    informe = 'inventario_valorizado'

    def test_trae_las_mismas_filas_que_existencia(self):
        """
        Valorizar un universo distinto haría que el total no cuadrara contra el
        informe de existencias.
        """
        self.assertEqual(
            self._ids(), {self.martillo.id, self.clavo.id, self.tornillo.id},
        )

    def test_trae_el_costo_y_la_valorizacion(self):
        fila = next(f for f in self._lista() if f['id'] == self.martillo.id)

        self.assertEqual(Decimal(fila['existencia']), Decimal('15'))
        self.assertEqual(Decimal(fila['costo_promedio']), Decimal('1000'))
        self.assertEqual(Decimal(fila['costo_total']), Decimal('15000'))
        self.assertEqual(Decimal(fila['costo']), Decimal('900'))
        self.assertEqual(Decimal(fila['precio']), Decimal('1500'))

    def test_el_costo_total_es_el_de_la_base_y_no_uno_recalculado(self):
        """
        `costo_total` lo mantiene `_afectar_inventario`; el informe lo muestra tal
        cual para no discrepar de la contabilidad. Si acá se recalculara, este
        item —al que se le desajustó el costo a mano— saldría distinto.
        """
        self.martillo.costo_total = Decimal('999')
        self.martillo.save(update_fields=['costo_total'])

        fila = next(f for f in self._lista() if f['id'] == self.martillo.id)
        self.assertEqual(Decimal(fila['costo_total']), Decimal('999'))

    def test_filtra_por_costo(self):
        self.assertEqual(
            self._ids(filtros=[
                {'propiedad': 'costo_promedio', 'operador': '>', 'valor': 100},
            ]),
            {self.martillo.id},
        )


class HistorialMovimientoTests(_InformeBase):
    """Las líneas de documento que movieron inventario."""

    informe = 'historial_movimiento'

    def test_trae_las_lineas_que_movieron_inventario_o_remision(self):
        self.assertEqual(
            self._ids(),
            {
                self.entrada.id,
                self.entrada_bodega.id,
                self.salida.id,
                self.despacho.id,
            },
        )

    def test_deja_por_fuera_lo_que_no_movio_ningun_saldo(self):
        """
        Las tres mitades de la invariante, una por línea: el documento sin
        aprobar no tocó un saldo, la línea sin almacén no llega al servicio y la
        que no declara operación no mueve nada.
        """
        ids = self._ids()

        self.assertNotIn(self.sin_aprobar.id, ids)
        self.assertNotIn(self.sin_almacen.id, ids)
        self.assertNotIn(self.sin_operacion.id, ids)

    def test_incluye_el_despacho_que_solo_mueve_remision(self):
        """
        Una remisión saca de disponible sin sacar de existencia: si la invariante
        mirara solo `operacion_inventario`, los despachos desaparecerían.
        """
        fila = next(f for f in self._lista() if f['id'] == self.despacho.id)

        self.assertEqual(fila['operacion_inventario'], 0)
        self.assertEqual(fila['operacion_remision'], -1)

    def test_la_cantidad_operada_reconstruye_la_existencia_del_item(self):
        """
        Es la relación entre este informe y los de saldo: el historial es cómo se
        llegó al saldo de hoy. La remisión no suma porque no mueve existencia.
        """
        movido = sum(
            Decimal(f['cantidad_operada']) for f in self._lista()
            if f['item_id'] == self.martillo.id and f['operacion_inventario'] != 0
        )

        self.assertEqual(movido, self.martillo.existencia)

    def test_trae_el_documento_y_sus_relaciones_resueltas(self):
        fila = next(f for f in self._lista() if f['id'] == self.salida.id)

        self.assertEqual(fila['documento_id'], self.salida.documento_id)
        self.assertEqual(fila['documento_fecha'], '2026-01-20')
        self.assertEqual(fila['documento_tipo_nombre'], 'Factura')
        self.assertEqual(fila['item_nombre'], 'Martillo')
        self.assertEqual(fila['almacen_nombre'], 'Principal')
        self.assertEqual(Decimal(fila['cantidad']), Decimal('3'))
        self.assertEqual(Decimal(fila['cantidad_operada']), Decimal('-3'))

    def test_ordena_del_movimiento_mas_reciente_al_mas_viejo(self):
        self.assertEqual(
            [f['documento_fecha'] for f in self._lista()],
            ['2026-01-21', '2026-01-20', '2026-01-11', '2026-01-10'],
        )

    def test_filtra_por_item_y_por_almacen(self):
        self.assertEqual(
            self._ids(filtros=[
                {'propiedad': 'almacen_id', 'operador': '=', 'valor': self.bodega.id},
            ]),
            {self.entrada_bodega.id},
        )
        self.assertEqual(
            self._ids(filtros=[
                {'propiedad': 'documento__fecha', 'operador': '>=', 'valor': '2026-01-20'},
            ]),
            {self.salida.id, self.despacho.id},
        )


class InformeContratoTests(_InformeBase):
    """Lo que comparten los cuatro informes: la selección del informe y el Excel."""

    informe = 'existencia'

    def test_informe_ausente_o_desconocido_es_400(self):
        request = APIRequestFactory().post('/inventario/informe/lista/', {}, format='json')
        self.assertEqual(
            _InformeViewSinPermisos.as_view({'post': 'lista'})(request).status_code, 400,
        )
        self.assertEqual(self._post('lista', informe='no_existe').status_code, 400)

    def test_un_filtro_de_otro_informe_es_400(self):
        """
        Cada informe declara sus `campos_filtrables` en su serializer, así que la
        whitelist cambia con el informe: `item__nombre` vale en el desglose por
        almacén y no en el consolidado, donde la columna se llama `nombre`.
        """
        self.assertEqual(
            self._post(
                'lista', informe='existencia',
                filtros=[{'propiedad': 'item__nombre', 'operador': '=', 'valor': 'x'}],
            ).status_code,
            400,
        )

    def _excel(self, informe):
        response = self._post('excel', informe=informe)
        self.assertEqual(response.status_code, 200)
        hoja = load_workbook(BytesIO(response.content)).active
        encabezados = [celda.value for celda in hoja[1]]
        ids = {hoja.cell(row=f, column=1).value for f in range(2, hoja.max_row + 1)}
        return encabezados, ids

    def test_cada_informe_exporta_su_propia_forma(self):
        """
        Cada informe declara su `exportar`, así que del mismo endpoint salen
        cuatro Excel distintos.
        """
        for informe, columna in (
            ('existencia', 'Existencia'),
            ('existencia_almacen', 'Almacén'),
            ('inventario_valorizado', 'Costo total'),
            ('historial_movimiento', 'Cantidad operada'),
        ):
            with self.subTest(informe=informe):
                encabezados, ids = self._excel(informe)

                self.assertIn(columna, encabezados)
                self.assertEqual(ids, self._ids(informe=informe))

    def test_el_excel_de_existencia_no_lleva_costo(self):
        encabezados, _ = self._excel('existencia')

        self.assertNotIn('Costo total', encabezados)
        self.assertNotIn('Almacén', encabezados)
