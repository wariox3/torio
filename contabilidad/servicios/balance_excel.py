"""
Excel de los informes de contabilidad.

No usa `ExportarExcelMixin`: el mixin escribe una tabla plana que arranca en la
fila de encabezados, y estos informes llevan encima un bloque con la empresa y el
corte, más un formato de número por columna. El mixin sigue sirviendo para todo
lo demás; generalizarlo para este caso lo convertiría en un motor de reportes.
"""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_FUENTE = 'Arial'
_FONDO = PatternFill('solid', fgColor='9CDFEB')

IMPORTE = '#,##0.00'
FECHA = 'yyyy-mm-dd'
TEXTO = 'General'

_FILA_ENCABEZADOS = 7

# (campo de la fila, encabezado, formato). El ancho no se declara: se calcula del
# contenido al escribir, como hacía itrio, así una razón social larga o un nombre
# de cuenta largo no quedan cortados en un tenant y sobrados en otro.
_CUENTA = (
    ('tipo', 'Tipo', TEXTO),
    ('codigo', 'Cuenta', TEXTO),
)
_CONTACTO = (
    ('nombre', 'Nombre Cuenta', TEXTO),
    ('identificacion', 'Identificación', TEXTO),
    ('contacto', 'Contacto', TEXTO),
)
_ASIENTO = (
    ('comprobante', 'Comprobante', TEXTO),
    ('numero', 'Numero', TEXTO),
    ('fecha', 'Fecha', FECHA),
)
_SALDOS = (
    ('saldo_anterior', 'Saldo anterior ($)', IMPORTE),
    ('debito', 'Debitos ($)', IMPORTE),
    ('credito', 'Creditos ($)', IMPORTE),
    ('saldo_final', 'Saldo actual ($)', IMPORTE),
)

# Una forma de tabla por informe. El encabezado de la columna del nombre cambia
# entre las dos familias —«Nombre de cuenta» en las de cuenta sola, «Nombre
# Cuenta» en las que abren por contacto—, así que las columnas se declaran
# enteras y no por composición ciega.
BALANCE = {
    'titulo': 'Balance de prueba',
    'hoja': 'balance_prueba',
    'columnas': _CUENTA + (('nombre', 'Nombre de cuenta', TEXTO),) + _SALDOS,
}
BALANCE_CONTACTO = {
    'titulo': 'Balance de prueba por contacto',
    'hoja': 'balance_prueba_contacto',
    'columnas': _CUENTA + _CONTACTO + _SALDOS,
}
AUXILIAR_CUENTA = {
    'titulo': 'Auxiliar cuenta',
    'hoja': 'auxiliar_cuenta',
    'columnas': _CUENTA + (('nombre', 'Nombre de cuenta', TEXTO),) + _SALDOS,
}
AUXILIAR_CONTACTO = {
    'titulo': 'Auxiliar por contacto',
    'hoja': 'auxiliar_contacto',
    'columnas': _CUENTA + _CONTACTO + _SALDOS,
}
AUXILIAR_GENERAL = {
    'titulo': 'Auxiliar general',
    'hoja': 'auxiliar_general',
    'columnas': _CUENTA + _CONTACTO + _ASIENTO + _SALDOS,
}
BASES = {
    'titulo': 'Informe de bases',
    'hoja': 'bases',
    'columnas': (
        ('codigo', 'Cuenta', TEXTO),
        ('nombre', 'Nombre Cuenta', TEXTO),
        ('identificacion', 'Identificación', TEXTO),
        ('contacto', 'Contacto', TEXTO),
        ('comprobante', 'Comprobante', TEXTO),
        ('numero', 'Numero', TEXTO),
        ('fecha', 'Fecha', FECHA),
        ('detalle', 'Detalle', TEXTO),
        ('debito', 'Debitos ($)', IMPORTE),
        ('credito', 'Creditos ($)', IMPORTE),
        ('base', 'Base ($)', IMPORTE),
    ),
}
CERTIFICADO_RETENCION = {
    'titulo': 'Certificado retenciones',
    # Único informe cuyo archivo no se llama como la hoja.
    'hoja': 'Certificado de retencion',
    'archivo': 'certificado_retencion',
    'columnas': (
        ('identificacion', 'Identificación', TEXTO),
        ('contacto', 'Contacto', TEXTO),
        ('codigo', 'Cuenta', TEXTO),
        ('nombre', 'Nombre cuenta', TEXTO),
        ('base_retenido', 'Monto del pago sujeto a retención ($)', IMPORTE),
        ('retenido', 'Retenido y consignado ($)', IMPORTE),
    ),
}
_ESTADO = (
    ('clase', 'Clase', TEXTO),
    ('grupo', 'Grupo', TEXTO),
    ('codigo', 'Codigo cuenta', TEXTO),
    ('nombre', 'Nombre cuenta', TEXTO),
    ('saldo', 'Saldo ($)', IMPORTE),
)
ESTADO_RESULTADOS = {
    'titulo': 'Estado de resultados',
    'hoja': 'estado_resultados',
    'columnas': _ESTADO,
}
ESTADO_SITUACION_FINANCIERA = {
    'titulo': 'Estado situacion financiera',
    'hoja': 'estado_situacion_financiera',
    'columnas': _ESTADO,
}


def excel(formato, filas, empresa, fecha_desde, fecha_hasta):
    """Devuelve los bytes del .xlsx de un informe, y el nombre del archivo."""
    columnas = formato['columnas']

    wb = Workbook()
    ws = wb.active
    ws.title = formato['hoja']

    anchos = [0] * len(columnas)
    _bloque_titulo(ws, formato['titulo'], empresa, fecha_desde, fecha_hasta, anchos)
    _encabezados(ws, columnas, anchos)

    for numero, fila in enumerate(filas, start=_FILA_ENCABEZADOS + 1):
        for indice, (campo, _, formato_numero) in enumerate(columnas):
            valor = fila[campo]
            celda = ws.cell(row=numero, column=indice + 1, value=valor)
            celda.font = Font(name=_FUENTE, size=10)
            celda.number_format = formato_numero
            anchos[indice] = max(anchos[indice], _largo(valor))

    for indice, ancho in enumerate(anchos):
        ws.column_dimensions[get_column_letter(indice + 1)].width = ancho + 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{formato.get('archivo') or formato['hoja']}.xlsx"


def _bloque_titulo(ws, titulo, empresa, fecha_desde, fecha_hasta, anchos):
    # El fondo va solo en la celda de arriba a la izquierda: es la que Excel
    # pinta a lo ancho de la combinación.
    ultima = get_column_letter(len(anchos))
    encabezado = (
        (titulo, 14, True, 'center'),
        (empresa, 12, True, 'center'),
        (f'Fecha generación: {timezone.localtime():%Y-%m-%d %H:%M}', 8, False, 'right'),
    )
    for numero, (texto, tamano, negrita, alineacion) in enumerate(encabezado, start=1):
        ws.merge_cells(f'A{numero}:{ultima}{numero}')
        celda = ws.cell(row=numero, column=1, value=texto)
        celda.font = Font(name=_FUENTE, size=tamano, bold=negrita)
        celda.alignment = Alignment(horizontal=alineacion)
        celda.fill = _FONDO
        anchos[0] = max(anchos[0], _largo(texto))

    rango = (f'Fecha desde: {fecha_desde}', f'Fecha hasta: {fecha_hasta}')
    for numero, texto in enumerate(rango, start=4):
        celda = ws.cell(row=numero, column=1, value=texto)
        celda.font = Font(name=_FUENTE, size=10)
        anchos[0] = max(anchos[0], _largo(texto))


def _encabezados(ws, columnas, anchos):
    for indice, (_, encabezado, _) in enumerate(columnas):
        celda = ws.cell(row=_FILA_ENCABEZADOS, column=indice + 1, value=encabezado)
        celda.font = Font(name=_FUENTE, size=10, bold=True)
        celda.fill = _FONDO
        anchos[indice] = max(anchos[indice], _largo(encabezado))


def _largo(valor):
    """Ancho que ocupa un valor, medido como lo escribe Python. Vacío no ocupa."""
    return 0 if valor is None else len(str(valor))
