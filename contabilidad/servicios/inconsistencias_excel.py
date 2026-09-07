"""
Excel de las inconsistencias de la contabilidad.

Tabla pelada: encabezados en la fila 1 y las claves que devuelve
`analizar_inconsistencias`, en su mismo orden. No lleva el bloque de título de
`balance_excel` —no es un informe con corte y razón social, es la lista de lo que
hay que corregir— ni pasa por `ExportarExcelMixin`, que escribe desde un
queryset serializado y acá las filas las arma el servicio.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_FUENTE = 'Arial'
_FONDO = PatternFill('solid', fgColor='9CDFEB')

# (clave de la inconsistencia, encabezado).
COLUMNAS = (
    ('comprobante_id', 'Comprobante'),
    ('comprobante_nombre', 'Nombre comprobante'),
    ('numero', 'Numero'),
    ('cuenta_id', 'Cuenta'),
    ('documento_id', 'Documento'),
    ('documento_tipo_nombre', 'Tipo de documento'),
    ('inconsistencia', 'Inconsistencia'),
)

ARCHIVO = 'inconsistencias.xlsx'


def excel(inconsistencias):
    """Devuelve los bytes del .xlsx de las inconsistencias, y el nombre del archivo."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'inconsistencias'

    # El ancho se calcula del contenido, como en `balance_excel`: el texto de la
    # inconsistencia es largo y variable, y una columna fija lo cortaría.
    anchos = [0] * len(COLUMNAS)
    for indice, (_, encabezado) in enumerate(COLUMNAS):
        celda = ws.cell(row=1, column=indice + 1, value=encabezado)
        celda.font = Font(name=_FUENTE, size=10, bold=True)
        celda.fill = _FONDO
        anchos[indice] = _largo(encabezado)

    for numero, inconsistencia in enumerate(inconsistencias, start=2):
        for indice, (clave, _) in enumerate(COLUMNAS):
            valor = inconsistencia[clave]
            celda = ws.cell(row=numero, column=indice + 1, value=valor)
            celda.font = Font(name=_FUENTE, size=10)
            anchos[indice] = max(anchos[indice], _largo(valor))

    for indice, ancho in enumerate(anchos):
        ws.column_dimensions[get_column_letter(indice + 1)].width = ancho + 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), ARCHIVO


def _largo(valor):
    """Ancho que ocupa un valor, medido como lo escribe Python. Vacío no ocupa."""
    return 0 if valor is None else len(str(valor))
