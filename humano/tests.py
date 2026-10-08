import io
from datetime import date
from decimal import Decimal

from django_tenants.test.cases import TenantTestCase
from rest_framework import permissions
from rest_framework.test import APIRequestFactory

from general.models import (
    GenCiudad,
    GenConfiguracion,
    GenContacto,
    GenEstado,
    GenIdentificacion,
    GenPais,
    GenTipoPersona,
)
from humano.models import (
    HumContrato,
    HumContratoTipo,
    HumGrupo,
    HumNovedad,
    HumPagoTipo,
    HumPeriodo,
    HumProgramacion,
    HumProgramacionDetalle,
)
from humano.serializers import HumProgramacionSerializer
from humano.servicios import ProgramacionError, cargar_contratos, eliminar_detalles
from humano.servicios import contrato as contrato_servicio
from humano.views.contrato import HumContratoViewSet
from seguridad.permissions import EsMiembroDelTenant, SuscripcionVigente
from utilidades.fechas import dias_prestacionales


class _ContratoViewSinPermisos(HumContratoViewSet):
    """Variante de la vista sin auth/permiso/throttle para probar los actions aislados."""
    authentication_classes = []
    permission_classes = [permissions.AllowAny]
    throttle_classes = []


class ValidacionContratoTests(TenantTestCase):
    """Un contacto no puede tener dos contratos abiertos ni contratos con fechas cruzadas."""

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        self.factory = APIRequestFactory()

        pais = GenPais.objects.create(id=1, nombre='Colombia')
        estado = GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais)
        ciudad = GenCiudad.objects.create(id=1, nombre='Bogotá', estado=estado)
        identificacion = GenIdentificacion.objects.create(id=1, nombre='CC')
        tipo_persona = GenTipoPersona.objects.create(id=1, nombre='Natural')
        self.contacto = GenContacto.objects.create(
            numero_identificacion='123',
            nombre_corto='Empleado',
            direccion='Calle 1',
            telefono='1',
            correo='e@e.com',
            identificacion=identificacion,
            ciudad=ciudad,
            tipo_persona=tipo_persona,
            empleado=True,
        )
        self.contacto2 = GenContacto.objects.create(
            numero_identificacion='456',
            nombre_corto='Otro',
            direccion='Calle 2',
            telefono='2',
            correo='o@o.com',
            identificacion=identificacion,
            ciudad=ciudad,
            tipo_persona=tipo_persona,
            empleado=True,
        )
        self.contrato_tipo = HumContratoTipo.objects.create(id=1, nombre='Fijo')
        self.grupo = HumGrupo.objects.create(nombre='Grupo 1')

    def _crear_contrato(self, **overrides):
        return HumContrato.objects.create(**{
            'fecha_desde': date(2026, 1, 1),
            'fecha_hasta': date(2026, 6, 30),
            'contrato_tipo': self.contrato_tipo,
            'contacto': self.contacto,
            'grupo': self.grupo,
            'estado_terminado': True,
            **overrides,
        })

    def _post(self, **overrides):
        view = _ContratoViewSinPermisos.as_view({'post': 'create'})
        payload = {
            'fecha_desde': '2026-07-01',
            'fecha_hasta': '2026-12-31',
            'contrato_tipo': self.contrato_tipo.id,
            'contacto': self.contacto.id,
            'grupo': self.grupo.id,
            **overrides,
        }
        return view(self.factory.post('/contrato/', payload, format='json'))

    def test_crea_contrato_sin_conflicto(self):
        response = self._post()

        self.assertEqual(response.status_code, 201)
        self.assertEqual(HumContrato.objects.count(), 1)

    def test_crea_contrato_despues_de_uno_terminado(self):
        # Terminado y sin cruce de fechas: se permite.
        self._crear_contrato()

        response = self._post()

        self.assertEqual(response.status_code, 201)
        self.assertEqual(HumContrato.objects.count(), 2)

    def test_rechaza_si_contacto_tiene_contrato_sin_terminar(self):
        self._crear_contrato(estado_terminado=False)

        response = self._post()

        self.assertEqual(response.status_code, 400)
        self.assertIn('detail', response.data)
        self.assertEqual(HumContrato.objects.count(), 1)

    def test_rechaza_fechas_cruzadas(self):
        # Terminado, pero el rango se solapa con el nuevo (2026-07-01 a 2026-12-31).
        self._crear_contrato(fecha_hasta=date(2026, 8, 31))

        response = self._post()

        self.assertEqual(response.status_code, 400)
        self.assertIn('detail', response.data)
        self.assertEqual(HumContrato.objects.count(), 1)

    def test_rechaza_cruce_por_un_solo_dia(self):
        self._crear_contrato(fecha_hasta=date(2026, 7, 1))

        response = self._post()

        self.assertEqual(response.status_code, 400)
        self.assertIn('detail', response.data)

    def test_permite_contrato_de_otro_contacto(self):
        # Mismo rango de fechas y sin terminar, pero de otro contacto.
        self._crear_contrato(estado_terminado=False, fecha_hasta=date(2026, 12, 31))

        response = self._post(contacto=self.contacto2.id)

        self.assertEqual(response.status_code, 201)
        self.assertEqual(HumContrato.objects.count(), 2)

    def test_rechaza_fecha_hasta_anterior_a_desde(self):
        response = self._post(fecha_desde='2026-12-31', fecha_hasta='2026-07-01')

        self.assertEqual(response.status_code, 400)
        self.assertIn('detail', response.data)

    def test_actualizar_no_choca_consigo_mismo(self):
        contrato = self._crear_contrato(estado_terminado=False)

        view = _ContratoViewSinPermisos.as_view({'patch': 'partial_update'})
        request = self.factory.patch('/contrato/', {'salario': 100}, format='json')
        response = view(request, pk=contrato.pk)

        self.assertEqual(response.status_code, 200)
        contrato.refresh_from_db()
        self.assertEqual(contrato.salario, 100)


    def test_resumen_cuenta_activos_y_terminados(self):
        self._crear_contrato()
        self._crear_contrato(contacto=self.contacto2, estado_terminado=False)
        self._crear_contrato(fecha_desde=date(2026, 7, 1), fecha_hasta=date(2026, 12, 31), estado_terminado=False)

        respuesta = _ContratoViewSinPermisos.as_view({'get': 'resumen'})(
            self.factory.get('/humano/contrato/resumen/'),
        )

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            {k: respuesta.data[k] for k in ('contratos', 'contratos_activos', 'contratos_terminados')},
            {'contratos': 3, 'contratos_activos': 2, 'contratos_terminados': 1},
        )

    def test_resumen_sin_contratos_responde_ceros(self):
        datos = contrato_servicio.resumen(hoy=date(2026, 10, 8))

        self.assertEqual(datos, {
            'fecha': date(2026, 10, 8), 'contratos': 0, 'contratos_activos': 0, 'contratos_terminados': 0,
            'ingresos_mes': 0, 'retiros_mes': 0,
        })

    def test_resumen_cuenta_ingresos_y_retiros_del_mes(self):
        hoy = date(2026, 10, 8)
        # Ingreso del mes, activo.
        self._crear_contrato(fecha_desde=date(2026, 10, 1), fecha_hasta=date(2026, 10, 1), estado_terminado=False)
        # Ingreso y retiro del mismo mes.
        self._crear_contrato(contacto=self.contacto2, fecha_desde=date(2026, 10, 2), fecha_hasta=date(2026, 10, 31))
        # Retiro del mes de un contrato que empezó antes.
        self._crear_contrato(fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 10, 5))
        # Termina en el mes pero sigue activo: un fijo con fecha de fin, no es retiro.
        self._crear_contrato(fecha_desde=date(2026, 4, 1), fecha_hasta=date(2026, 10, 20), estado_terminado=False)
        # Fuera del mes: septiembre del mismo año y octubre de otro año.
        self._crear_contrato(fecha_desde=date(2026, 9, 30), fecha_hasta=date(2026, 9, 30))
        self._crear_contrato(fecha_desde=date(2025, 10, 1), fecha_hasta=date(2025, 10, 31))

        datos = contrato_servicio.resumen(hoy=hoy)

        self.assertEqual(datos['fecha'], hoy)
        self.assertEqual(datos['ingresos_mes'], 2)
        self.assertEqual(datos['retiros_mes'], 2)
        self.assertEqual(datos['contratos'], 6)

    def test_resumen_no_exige_el_permiso_del_modelo(self):
        """Es el tablero de inicio: lo ve cualquier miembro con suscripción, como `cartera-resumen`."""
        self.assertEqual(
            HumContratoViewSet.resumen.kwargs['permission_classes'], [EsMiembroDelTenant, SuscripcionVigente],
        )

class CargarContratosTests(TenantTestCase):
    """`cargar_contratos`: un detalle por contrato del grupo según el tipo de pago."""

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        pais = GenPais.objects.create(id=1, nombre='Colombia')
        ciudad = GenCiudad.objects.create(
            id=1, nombre='Bogotá', estado=GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais),
        )
        self.contacto = GenContacto.objects.create(
            numero_identificacion='123', nombre_corto='Empleado', direccion='Calle 1', telefono='1',
            correo='e@e.com', identificacion=GenIdentificacion.objects.create(id=1, nombre='CC'),
            ciudad=ciudad, tipo_persona=GenTipoPersona.objects.create(id=1, nombre='Natural'), empleado=True,
        )
        GenConfiguracion.objects.update_or_create(pk=1, defaults={
            'hum_factor': Decimal('8'),
            'hum_salario_minimo': Decimal('1423500'),
            'hum_auxilio_transporte': Decimal('200000'),
        })
        self.indefinido = HumContratoTipo.objects.create(id=1, nombre='Indefinido')
        self.fijo = HumContratoTipo.objects.create(id=2, nombre='Fijo')
        self.practica = HumContratoTipo.objects.create(id=5, nombre='Practica estudiantil')
        for id, nombre in ((1, 'Nomina'), (2, 'Prima'), (3, 'Cesantia')):
            HumPagoTipo.objects.create(id=id, nombre=nombre)
        self.grupo = HumGrupo.objects.create(nombre='Grupo 1')

    def _programacion(self, **overrides):
        return HumProgramacion.objects.create(**{
            'fecha_desde': date(2026, 3, 1), 'fecha_hasta': date(2026, 3, 15),
            'fecha_hasta_periodo': date(2026, 3, 15), 'dias_reales': 15,
            'grupo': self.grupo, 'pago_tipo_id': 1, **overrides,
        })

    def _contrato(self, **overrides):
        return HumContrato.objects.create(**{
            'fecha_desde': date(2026, 1, 1), 'fecha_hasta': date(2026, 1, 1),
            'salario': Decimal('2000000'), 'contrato_tipo': self.indefinido,
            'contacto': self.contacto, 'grupo': self.grupo, **overrides,
        })

    # ---- nómina ----

    def test_nomina_indefinido_quincena_completa(self):
        contrato = self._contrato()
        programacion = self._programacion()

        self.assertEqual(cargar_contratos(programacion), 1)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertEqual(detalle.contrato_id, contrato.id)
        self.assertEqual((detalle.fecha_desde, detalle.fecha_hasta), (date(2026, 3, 1), date(2026, 3, 15)))
        self.assertEqual(detalle.dias, 15)
        self.assertEqual(detalle.diurna, 120)
        self.assertFalse(detalle.ingreso)
        programacion.refresh_from_db()
        self.assertEqual(programacion.contratos, 1)

    def test_nomina_ingreso_en_el_periodo(self):
        self._contrato(fecha_desde=date(2026, 3, 6))
        programacion = self._programacion()

        cargar_contratos(programacion)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertTrue(detalle.ingreso)
        self.assertEqual(detalle.dias, 10)

    def test_nomina_resta_dias_de_novedad(self):
        contrato = self._contrato()
        HumNovedad.objects.create(contrato=contrato, fecha_desde=date(2026, 3, 10), fecha_hasta=date(2026, 3, 20))
        programacion = self._programacion()

        cargar_contratos(programacion)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertEqual(detalle.dias_novedad, 6)
        self.assertEqual(detalle.dias, 9)

    def test_nomina_practica_no_cotiza_ni_recibe_transporte(self):
        self._contrato(contrato_tipo=self.practica, fecha_hasta=date(2026, 12, 31))
        programacion = self._programacion()

        cargar_contratos(programacion)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertFalse(detalle.descuento_salud)
        self.assertFalse(detalle.descuento_pension)
        self.assertFalse(detalle.pago_auxilio_transporte)

    def test_nomina_contrato_fijo_vencido_sin_terminar_marca_error(self):
        self._contrato(contrato_tipo=self.fijo, fecha_hasta=date(2026, 2, 28))
        programacion = self._programacion()

        cargar_contratos(programacion)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertTrue(detalle.error_terminacion)
        self.assertEqual(detalle.dias, 0)

    def test_no_duplica_contratos_ya_cargados(self):
        self._contrato()
        programacion = self._programacion()
        cargar_contratos(programacion)

        self.assertEqual(cargar_contratos(programacion), 0)
        self.assertEqual(HumProgramacionDetalle.objects.filter(programacion=programacion).count(), 1)

    def test_programacion_generada_no_carga(self):
        self._contrato()
        programacion = self._programacion(estado_generado=True)

        with self.assertRaises(ProgramacionError):
            cargar_contratos(programacion)
        self.assertFalse(HumProgramacionDetalle.objects.exists())

    def test_contratos_de_otro_grupo_no_entran(self):
        self._contrato(grupo=HumGrupo.objects.create(nombre='Otro'))
        self.assertEqual(cargar_contratos(self._programacion()), 0)

    # ---- prima y cesantía ----

    def test_prima_suma_auxilio_al_salario_promedio(self):
        self._contrato(auxilio_transporte=True)
        programacion = self._programacion(
            pago_tipo_id=2, fecha_desde=date(2026, 1, 1),
            fecha_hasta=date(2026, 6, 30), fecha_hasta_periodo=date(2026, 6, 30),
        )

        cargar_contratos(programacion)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertEqual(detalle.dias, 180)
        self.assertEqual(detalle.salario_promedio, Decimal('2200000'))

    def test_cesantia_sin_nomina_toma_el_salario_como_minimo(self):
        self._contrato()
        programacion = self._programacion(
            pago_tipo_id=3, fecha_desde=date(2026, 1, 1),
            fecha_hasta=date(2026, 6, 30), fecha_hasta_periodo=date(2026, 6, 30),
        )

        cargar_contratos(programacion)

        detalle = HumProgramacionDetalle.objects.get(programacion=programacion)
        self.assertEqual(detalle.dias, 180)
        self.assertEqual(detalle.base_prestacion, 0)
        self.assertEqual(detalle.salario_promedio, Decimal('2000000'))

    def test_dias_prestacionales(self):
        self.assertEqual(dias_prestacionales(date(2026, 1, 1), date(2026, 6, 30)), 180)
        self.assertEqual(dias_prestacionales(date(2026, 1, 1), date(2026, 2, 28)), 60)
        self.assertEqual(dias_prestacionales(date(2026, 7, 1), date(2026, 12, 30)), 180)
        # 30/360: el 31 cuenta como 30.
        self.assertEqual(dias_prestacionales(date(2026, 1, 1), date(2026, 12, 31)), 360)
        self.assertEqual(dias_prestacionales(date(2026, 7, 1), date(2026, 12, 31)), 180)
        self.assertEqual(dias_prestacionales(date(2026, 1, 31), date(2026, 3, 31)), 61)
        # Febrero completa 30 solo en su último día: en bisiesto el 28 no lo es.
        self.assertEqual(dias_prestacionales(date(2028, 1, 1), date(2028, 2, 29)), 60)
        self.assertEqual(dias_prestacionales(date(2028, 1, 1), date(2028, 2, 28)), 58)

    # ---- eliminar detalles ----

    def _cargada(self, **overrides):
        self._contrato()
        self._contrato(contacto=GenContacto.objects.create(
            numero_identificacion='456', nombre_corto='Otro', direccion='Calle 2', telefono='2',
            correo='o@o.com', identificacion_id=1, ciudad_id=1, tipo_persona_id=1, empleado=True,
        ))
        programacion = self._programacion(**overrides)
        cargar_contratos(programacion)
        programacion.refresh_from_db()
        return programacion

    def test_eliminar_todos_los_detalles(self):
        programacion = self._cargada()

        self.assertEqual(eliminar_detalles(programacion), 2)
        programacion.refresh_from_db()
        self.assertEqual(programacion.contratos, 0)

    def test_eliminar_solo_los_indicados(self):
        programacion = self._cargada()
        detalle = HumProgramacionDetalle.objects.filter(programacion=programacion).first()

        self.assertEqual(eliminar_detalles(programacion, [detalle.id]), 1)
        programacion.refresh_from_db()
        self.assertEqual(programacion.contratos, 1)

    def test_eliminar_ignora_detalles_de_otra_programacion(self):
        programacion = self._cargada()
        otra = self._programacion()
        cargar_contratos(otra)
        ajeno = HumProgramacionDetalle.objects.filter(programacion=otra).first()

        self.assertEqual(eliminar_detalles(programacion, [ajeno.id]), 0)
        self.assertTrue(HumProgramacionDetalle.objects.filter(pk=ajeno.pk).exists())

    def test_no_elimina_en_programacion_generada_o_aprobada(self):
        programacion = self._cargada()
        for estado in ('estado_generado', 'estado_aprobado'):
            HumProgramacion.objects.filter(pk=programacion.pk).update(
                estado_generado=estado == 'estado_generado', estado_aprobado=estado == 'estado_aprobado',
            )

            with self.assertRaises(ProgramacionError):
                eliminar_detalles(programacion)
            self.assertEqual(HumProgramacionDetalle.objects.filter(programacion=programacion).count(), 2)


class PeriodoProgramacionTests(TenantTestCase):
    """`fecha_hasta_periodo`, `dias` y `dias_reales` los calcula el servidor."""

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        self.grupo = HumGrupo.objects.create(nombre='Grupo 1')
        HumPagoTipo.objects.create(id=1, nombre='Nomina')

    def _crear(self, fecha_desde, fecha_hasta, **extra):
        serializer = HumProgramacionSerializer(data={
            'fecha_desde': fecha_desde, 'fecha_hasta': fecha_hasta,
            'grupo': self.grupo.id, 'pago_tipo': 1, **extra,
        })
        serializer.is_valid(raise_exception=True)
        return serializer.save()

    def test_segunda_quincena_de_mes_de_31_cierra_el_31(self):
        programacion = self._crear('2026-03-16', '2026-03-30')

        self.assertEqual(programacion.fecha_hasta_periodo, date(2026, 3, 31))
        self.assertEqual(programacion.dias, 15)
        self.assertEqual(programacion.dias_reales, 15)

    def test_mes_de_30_cierra_en_la_fecha_hasta(self):
        programacion = self._crear('2026-04-16', '2026-04-30')

        self.assertEqual(programacion.fecha_hasta_periodo, date(2026, 4, 30))

    def test_febrero_completo_son_30_dias(self):
        programacion = self._crear('2026-02-01', '2026-02-28')

        self.assertEqual(programacion.fecha_hasta_periodo, date(2026, 2, 28))
        self.assertEqual(programacion.dias, 30)
        self.assertEqual(programacion.dias_reales, 28)

    def test_lo_que_mande_el_usuario_se_ignora(self):
        programacion = self._crear(
            '2026-03-01', '2026-03-15', fecha_hasta_periodo='2026-12-31', dias=99, dias_reales=99,
        )

        self.assertEqual(programacion.fecha_hasta_periodo, date(2026, 3, 15))
        self.assertEqual((programacion.dias, programacion.dias_reales), (15, 15))

    def test_editar_las_fechas_recalcula(self):
        programacion = self._crear('2026-03-01', '2026-03-15')

        serializer = HumProgramacionSerializer(programacion, data={'fecha_hasta': '2026-03-30'}, partial=True)
        serializer.is_valid(raise_exception=True)
        programacion = serializer.save()

        self.assertEqual(programacion.fecha_hasta_periodo, date(2026, 3, 31))
        self.assertEqual(programacion.dias, 30)

    def test_el_periodo_se_guarda(self):
        # `HumProgramacion.periodo` es el FK: un método con ese nombre lo tapaba y
        # el campo dejaba de existir para Django.
        HumPeriodo.objects.create(id=1, codigo='Q', nombre='Quincenal', dias=15)

        programacion = self._crear('2026-03-01', '2026-03-15', periodo=1)

        programacion.refresh_from_db()
        self.assertEqual(programacion.periodo_id, 1)

    def test_sin_periodo_toma_el_del_grupo(self):
        self.grupo.periodo = HumPeriodo.objects.create(id=1, codigo='Q', nombre='Quincenal', dias=15)
        self.grupo.save()

        self.assertEqual(self._crear('2026-03-01', '2026-03-15').periodo_id, 1)
        self.assertEqual(self._crear('2026-03-01', '2026-03-15', periodo=None).periodo_id, 1)

    def test_el_periodo_enviado_manda_sobre_el_del_grupo(self):
        self.grupo.periodo = HumPeriodo.objects.create(id=1, codigo='Q', nombre='Quincenal', dias=15)
        self.grupo.save()
        HumPeriodo.objects.create(id=2, codigo='M', nombre='Mensual', dias=30)

        self.assertEqual(self._crear('2026-03-01', '2026-03-30', periodo=2).periodo_id, 2)

    def test_grupo_sin_periodo_queda_sin_periodo(self):
        self.assertIsNone(self._crear('2026-03-01', '2026-03-15').periodo_id)

    def test_fecha_hasta_anterior_a_desde(self):
        serializer = HumProgramacionSerializer(data={
            'fecha_desde': '2026-03-15', 'fecha_hasta': '2026-03-01', 'grupo': self.grupo.id, 'pago_tipo': 1,
        })
        self.assertFalse(serializer.is_valid())


class EditarProgramacionDetalleTests(TenantTestCase):
    """PUT/PATCH de `programacion-detalle`: solo con la programación sin generar."""

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        from humano.views.programacion_detalle import HumProgramacionDetalleViewSet

        class Vista(HumProgramacionDetalleViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self.vista = Vista.as_view({'patch': 'partial_update', 'put': 'update'})
        self.crear = Vista.as_view({'post': 'create'})
        self.factory = APIRequestFactory()

        pais = GenPais.objects.create(id=1, nombre='Colombia')
        ciudad = GenCiudad.objects.create(
            id=1, nombre='Bogotá', estado=GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais),
        )
        identificacion = GenIdentificacion.objects.create(id=1, nombre='CC')
        tipo_persona = GenTipoPersona.objects.create(id=1, nombre='Natural')
        contactos = [
            GenContacto.objects.create(
                numero_identificacion=numero, nombre_corto=numero, direccion='x', telefono='1', correo='e@e.com',
                identificacion=identificacion, ciudad=ciudad, tipo_persona=tipo_persona, empleado=True,
            )
            for numero in ('1', '2')
        ]
        grupo = HumGrupo.objects.create(nombre='Grupo 1')
        HumPagoTipo.objects.create(id=1, nombre='Nomina')
        tipo = HumContratoTipo.objects.create(id=1, nombre='Indefinido')
        self.contratos = [
            HumContrato.objects.create(
                fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 1), contrato_tipo=tipo,
                contacto=contacto, grupo=grupo,
            )
            for contacto in contactos
        ]
        self.programacion = HumProgramacion.objects.create(
            fecha_desde=date(2026, 3, 1), fecha_hasta=date(2026, 3, 15), fecha_hasta_periodo=date(2026, 3, 15),
            grupo=grupo, pago_tipo_id=1,
        )
        self.detalle = HumProgramacionDetalle.objects.create(programacion=self.programacion, contrato=self.contratos[0])

    def _patch(self, datos):
        return self.vista(self.factory.patch('/', datos, format='json'), pk=self.detalle.pk)

    def test_edita_las_horas(self):
        respuesta = self._patch({'extra_diurna': 4, 'pago_auxilio_transporte': False})

        self.assertEqual(respuesta.status_code, 200)
        self.detalle.refresh_from_db()
        self.assertEqual(self.detalle.extra_diurna, 4)
        self.assertFalse(self.detalle.pago_auxilio_transporte)

    def test_put_completo(self):
        datos = {'programacion': self.programacion.id, 'contrato': self.contratos[0].id, 'diurna': 80}

        respuesta = self.vista(self.factory.put('/', datos, format='json'), pk=self.detalle.pk)

        self.assertEqual(respuesta.status_code, 200)
        self.detalle.refresh_from_db()
        self.assertEqual(self.detalle.diurna, 80)

    def test_no_cambia_de_contrato(self):
        respuesta = self._patch({'contrato': self.contratos[1].id})

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('detail', respuesta.data)
        self.detalle.refresh_from_db()
        self.assertEqual(self.detalle.contrato_id, self.contratos[0].id)

    def test_no_edita_con_la_programacion_generada(self):
        HumProgramacion.objects.filter(pk=self.programacion.pk).update(estado_generado=True)

        respuesta = self._patch({'extra_diurna': 4})

        self.assertEqual(respuesta.status_code, 400)
        self.detalle.refresh_from_db()
        self.assertEqual(self.detalle.extra_diurna, 0)

    def test_crea_con_la_programacion_sin_generar(self):
        datos = {'programacion': self.programacion.id, 'contrato': self.contratos[1].id}

        respuesta = self.crear(self.factory.post('/', datos, format='json'))

        self.assertEqual(respuesta.status_code, 201)

    def test_no_crea_con_la_programacion_generada(self):
        HumProgramacion.objects.filter(pk=self.programacion.pk).update(estado_generado=True)
        datos = {'programacion': self.programacion.id, 'contrato': self.contratos[1].id}

        respuesta = self.crear(self.factory.post('/', datos, format='json'))

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('detail', respuesta.data)
        self.assertEqual(HumProgramacionDetalle.objects.count(), 1)

    def test_el_backend_tampoco_crea_con_la_programacion_generada(self):
        from django.core.exceptions import ValidationError

        HumProgramacion.objects.filter(pk=self.programacion.pk).update(estado_generado=True)

        with self.assertRaises(ValidationError):
            HumProgramacionDetalle.objects.create(programacion=self.programacion, contrato=self.contratos[1])
        with self.assertRaises(ValidationError):
            HumProgramacionDetalle.objects.bulk_create([
                HumProgramacionDetalle(programacion=self.programacion, contrato=self.contratos[1]),
            ])
        self.assertEqual(HumProgramacionDetalle.objects.count(), 1)


class ImportarHorasProgramacionDetalleTests(TenantTestCase):
    """
    `programacion-detalle/importar-horas/`: actualiza las horas de los detalles de
    una programación abierta, todo o nada.
    """

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        from humano.views.programacion_detalle import HumProgramacionDetalleViewSet

        class Vista(HumProgramacionDetalleViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self.importar = Vista.as_view({'post': 'importar'})
        self.plantilla = Vista.as_view({'get': 'importar_ejemplo'})
        self.factory = APIRequestFactory()

        pais = GenPais.objects.create(id=1, nombre='Colombia')
        ciudad = GenCiudad.objects.create(
            id=1, nombre='Bogotá', estado=GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais),
        )
        identificacion = GenIdentificacion.objects.create(id=1, nombre='CC')
        tipo_persona = GenTipoPersona.objects.create(id=1, nombre='Natural')
        grupo = HumGrupo.objects.create(nombre='Grupo 1')
        HumPagoTipo.objects.create(id=1, nombre='Nomina')
        tipo = HumContratoTipo.objects.create(id=1, nombre='Indefinido')
        contratos = [
            HumContrato.objects.create(
                fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 1), contrato_tipo=tipo, grupo=grupo,
                contacto=GenContacto.objects.create(
                    numero_identificacion=numero, nombre_corto=nombre, direccion='x', telefono='1',
                    correo='e@e.com', identificacion=identificacion, ciudad=ciudad,
                    tipo_persona=tipo_persona, empleado=True,
                ),
            )
            for numero, nombre in (('1', 'Beatriz'), ('2', 'Andrés'))
        ]
        self.programacion, self.otra = (
            HumProgramacion.objects.create(
                fecha_desde=date(2026, 3, 1), fecha_hasta=date(2026, 3, 15),
                fecha_hasta_periodo=date(2026, 3, 15), grupo=grupo, pago_tipo_id=1,
            )
            for _ in range(2)
        )
        self.detalles = [
            HumProgramacionDetalle.objects.create(programacion=self.programacion, contrato=contrato)
            for contrato in contratos
        ]
        self.ajeno = HumProgramacionDetalle.objects.create(programacion=self.otra, contrato=contratos[0])

    def _archivo(self, filas, encabezados=None):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        from humano.serializers import HumProgramacionDetalleImportarHorasSerializer
        from utilidades.mixins import ImportarExcelMixin

        serializer = HumProgramacionDetalleImportarHorasSerializer(self.programacion)
        wb = Workbook()
        ws = wb.active
        ws.append(encabezados or [
            ImportarExcelMixin._encabezado_importar(campo, encabezado, serializer.campos_requeridos)
            for campo, encabezado in serializer.campos_excel
        ])
        for fila in filas:
            ws.append(fila)
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile(
            'horas.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def _post(self, filas, programacion=None, **kwargs):
        datos = {
            'archivo': self._archivo(filas, **kwargs),
            'programacion_id': (programacion or self.programacion).id,
        }
        return self.importar(self.factory.post('/', datos, format='multipart'))

    @staticmethod
    def _fila(detalle, *horas):
        """ID, identificación, nombre y las 11 horas (las que falten van en 0)."""
        horas = list(horas) + [0] * (11 - len(horas))
        return [detalle.id, 'x', 'x', *horas]

    def test_actualiza_las_horas(self):
        respuesta = self._post([
            self._fila(self.detalles[0], 80, 8, 0, 0, 2.5),
            self._fila(self.detalles[1], 96),
        ])

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual(respuesta.data, {'creados': 2})
        self.detalles[0].refresh_from_db()
        self.assertEqual(self.detalles[0].diurna, 80)
        self.assertEqual(self.detalles[0].nocturna, 8)
        self.assertEqual(self.detalles[0].extra_diurna, Decimal('2.5'))
        self.detalles[1].refresh_from_db()
        self.assertEqual(self.detalles[1].diurna, 96)

    def test_celda_vacia_es_cero(self):
        HumProgramacionDetalle.objects.filter(pk=self.detalles[0].pk).update(nocturna=5)
        fila = self._fila(self.detalles[0], 80)
        fila[4] = None  # nocturna

        respuesta = self._post([fila])

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.detalles[0].refresh_from_db()
        self.assertEqual(self.detalles[0].nocturna, 0)

    def test_no_importa_con_la_programacion_generada(self):
        HumProgramacion.objects.filter(pk=self.programacion.pk).update(estado_generado=True)

        respuesta = self._post([self._fila(self.detalles[0], 80)])

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('detail', respuesta.data)
        self.detalles[0].refresh_from_db()
        self.assertEqual(self.detalles[0].diurna, 0)

    def test_programacion_inexistente_da_404(self):
        respuesta = self.importar(self.factory.post(
            '/', {'archivo': self._archivo([]), 'programacion_id': 999999}, format='multipart',
        ))

        self.assertEqual(respuesta.status_code, 404)

    def test_detalle_de_otra_programacion_revierte_todo(self):
        respuesta = self._post([
            self._fila(self.detalles[0], 80),
            self._fila(self.ajeno, 80),
        ])

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['fase'], 'negocio')
        self.assertEqual(respuesta.data['errores'][0]['fila'], 3)
        self.detalles[0].refresh_from_db()
        self.assertEqual(self.detalles[0].diurna, 0)

    def test_detalle_repetido(self):
        respuesta = self._post([
            self._fila(self.detalles[0], 80),
            self._fila(self.detalles[0], 90),
        ])

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['errores'], [
            {'fila': 3, 'mensaje': f'El detalle {self.detalles[0].id} ya viene en la fila 2.'},
        ])

    def test_horas_negativas(self):
        respuesta = self._post([self._fila(self.detalles[0], -1)])

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['fase'], 'negocio')
        self.assertIn('Diurna no puede ser negativo', respuesta.data['errores'][0]['mensaje'])

    def test_horas_que_no_son_numero(self):
        respuesta = self._post([self._fila(self.detalles[0], 'ocho')])

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('Diurna', respuesta.data['errores'][0]['mensaje'])

    def test_encabezados_que_no_coinciden(self):
        respuesta = self._post([[self.detalles[0].id, 80]], encabezados=['ID *', 'Horas'])

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['fase'], 'encabezados')

    def test_la_plantilla_trae_los_detalles_de_la_programacion(self):
        from openpyxl import load_workbook

        HumProgramacionDetalle.objects.filter(pk=self.detalles[0].pk).update(diurna=80)

        respuesta = self.plantilla(self.factory.get('/', {'programacion_id': self.programacion.id}))

        self.assertEqual(respuesta.status_code, 200)
        filas = list(load_workbook(io.BytesIO(respuesta.content)).active.iter_rows(values_only=True))
        self.assertEqual(filas[0][:4], ('ID *', 'Identificación', 'Nombre', 'Diurna'))
        # Ordenados por nombre y sin el detalle de la otra programación.
        self.assertEqual([f[:4] for f in filas[1:]], [
            (self.detalles[1].id, '2', 'Andrés', 0),
            (self.detalles[0].id, '1', 'Beatriz', 80),
        ])

    def test_la_plantilla_sube_tal_cual(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        respuesta = self.plantilla(self.factory.get('/', {'programacion_id': self.programacion.id}))

        archivo = SimpleUploadedFile('horas.xlsx', respuesta.content)
        respuesta = self.importar(self.factory.post(
            '/', {'archivo': archivo, 'programacion_id': self.programacion.id}, format='multipart',
        ))

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual(respuesta.data, {'creados': 2})


class ExportarProgramacionDetalleTests(TenantTestCase):
    """`programacion-detalle/excel/`: el exportar convencional, con los filtros de `lista`."""

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        from humano.views.programacion_detalle import HumProgramacionDetalleViewSet

        class Vista(HumProgramacionDetalleViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self.excel = Vista.as_view({'post': 'excel'})
        self.factory = APIRequestFactory()

        pais = GenPais.objects.create(id=1, nombre='Colombia')
        ciudad = GenCiudad.objects.create(
            id=1, nombre='Bogotá', estado=GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais),
        )
        grupo = HumGrupo.objects.create(nombre='Grupo 1')
        HumPagoTipo.objects.create(id=1, nombre='Nomina')
        contrato = HumContrato.objects.create(
            fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 1), grupo=grupo,
            contrato_tipo=HumContratoTipo.objects.create(id=1, nombre='Indefinido'),
            contacto=GenContacto.objects.create(
                numero_identificacion='123', nombre_corto='Beatriz', direccion='x', telefono='1',
                correo='e@e.com', ciudad=ciudad, empleado=True,
                identificacion=GenIdentificacion.objects.create(id=1, nombre='CC'),
                tipo_persona=GenTipoPersona.objects.create(id=1, nombre='Natural'),
            ),
        )
        self.programacion, otra = (
            HumProgramacion.objects.create(
                fecha_desde=date(2026, 3, 1), fecha_hasta=date(2026, 3, 15),
                fecha_hasta_periodo=date(2026, 3, 15), grupo=grupo, pago_tipo_id=1,
            )
            for _ in range(2)
        )
        self.detalle = HumProgramacionDetalle.objects.create(
            programacion=self.programacion, contrato=contrato, diurna=80, ingreso=True,
        )
        HumProgramacionDetalle.objects.create(programacion=otra, contrato=contrato)

    def test_exporta_los_detalles_de_la_programacion_filtrada(self):
        from openpyxl import load_workbook

        filtros = [{'propiedad': 'programacion_id', 'operador': '=', 'valor': self.programacion.id}]

        respuesta = self.excel(self.factory.post('/', {'filtros': filtros}, format='json'))

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('programacion_detalles.xlsx', respuesta['Content-Disposition'])
        filas = list(load_workbook(io.BytesIO(respuesta.content)).active.iter_rows(values_only=True))
        encabezados = filas[0]
        self.assertEqual(encabezados[:5], ('ID', 'Programación', 'Contrato', 'Identificación', 'Empleado'))
        self.assertEqual(len(filas), 2)
        fila = dict(zip(encabezados, filas[1], strict=True))
        self.assertEqual(fila['ID'], self.detalle.id)
        self.assertEqual(fila['Identificación'], '123')
        self.assertEqual(fila['Empleado'], 'Beatriz')
        self.assertEqual(fila['Diurna'], 80)
        self.assertEqual(fila['Ingreso'], 'Sí')
        self.assertEqual(fila['Retiro'], 'No')


class InformesNominaTests(TenantTestCase):
    """
    Informes `nomina` (`documento-informe`) y `nomina_detalle`
    (`documento-detalle-informe`): todas las nóminas, aprobadas o no, y nada más,
    filtrables por programación.
    """

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        from general.models import (
            GenDocumento,
            GenDocumentoClase,
            GenDocumentoDetalle,
            GenDocumentoTipo,
        )
        from general.views.documento_detalle_informe import (
            GenDocumentoDetalleInformeViewSet,
        )
        from general.views.documento_informe import GenDocumentoInformeViewSet
        from humano.models import HumConcepto

        class Vista(GenDocumentoDetalleInformeViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        class VistaDocumento(GenDocumentoInformeViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self.lista = Vista.as_view({'post': 'lista'})
        self.excel = Vista.as_view({'post': 'excel'})
        self.lista_documento = VistaDocumento.as_view({'post': 'lista'})
        self.excel_documento = VistaDocumento.as_view({'post': 'excel'})
        self.factory = APIRequestFactory()

        pais = GenPais.objects.create(id=1, nombre='Colombia')
        ciudad = GenCiudad.objects.create(
            id=1, nombre='Bogotá', estado=GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais),
        )
        contacto = GenContacto.objects.create(
            numero_identificacion='123', nombre_corto='Beatriz', direccion='x', telefono='1',
            correo='e@e.com', ciudad=ciudad, empleado=True,
            identificacion=GenIdentificacion.objects.create(id=1, nombre='CC'),
            tipo_persona=GenTipoPersona.objects.create(id=1, nombre='Natural'),
        )
        grupo = HumGrupo.objects.create(nombre='Grupo 1')
        HumPagoTipo.objects.create(id=1, nombre='Nomina')
        contrato = HumContrato.objects.create(
            fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 1), grupo=grupo, contacto=contacto,
            contrato_tipo=HumContratoTipo.objects.create(id=1, nombre='Indefinido'),
        )
        self.programacion, otra = (
            HumProgramacion.objects.create(
                fecha_desde=date(2026, 3, 1), fecha_hasta=date(2026, 3, 15),
                fecha_hasta_periodo=date(2026, 3, 15), grupo=grupo, pago_tipo_id=1,
            )
            for _ in range(2)
        )
        nomina = GenDocumentoTipo.objects.create(
            id=14, nombre='NOMINA', documento_clase=GenDocumentoClase.objects.create(id=701, nombre='Nomina'),
        )
        factura = GenDocumentoTipo.objects.create(
            id=1, nombre='FACTURA', documento_clase=GenDocumentoClase.objects.create(id=100, nombre='Factura'),
        )
        concepto = HumConcepto.objects.create(id=1, nombre='SALARIO')

        def detalle(tipo, programacion=None, **campos):
            documento = GenDocumento.objects.create(
                documento_tipo=tipo, fecha=date(2026, 3, 1), fecha_hasta=date(2026, 3, 15),
                contacto=contacto, contrato=contrato,
                programacion_detalle=HumProgramacionDetalle.objects.create(
                    programacion=programacion, contrato=contrato,
                ) if programacion else None,
            )
            return GenDocumentoDetalle.objects.create(documento=documento, concepto=concepto, **campos)

        self.de_la_programacion = detalle(nomina, self.programacion, devengado=1000, dias=15)
        self.de_otra = detalle(nomina, otra)
        detalle(factura)

    def _post(self, vista, filtros=None, informe='nomina_detalle'):
        datos = {'informe': informe, 'filtros': filtros or []}
        return vista(self.factory.post('/', datos, format='json'))

    def test_solo_trae_detalles_de_nomina(self):
        respuesta = self._post(self.lista)

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        ids = {fila['id'] for fila in respuesta.data['results']}
        self.assertEqual(ids, {self.de_la_programacion.id, self.de_otra.id})

    def test_filtra_por_programacion(self):
        filtros = [{
            'propiedad': 'documento__programacion_detalle__programacion_id',
            'operador': '=', 'valor': self.programacion.id,
        }]

        respuesta = self._post(self.lista, filtros)

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual([fila['id'] for fila in respuesta.data['results']], [self.de_la_programacion.id])
        fila = respuesta.data['results'][0]
        self.assertEqual(fila['contacto_numero_identificacion'], '123')
        self.assertEqual(fila['concepto_nombre'], 'SALARIO')

    def test_excel_con_las_columnas_del_informe(self):
        from openpyxl import load_workbook

        filtros = [{
            'propiedad': 'documento__programacion_detalle__programacion_id',
            'operador': '=', 'valor': self.programacion.id,
        }]

        respuesta = self._post(self.excel, filtros)

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('nomina_detalles.xlsx', respuesta['Content-Disposition'])
        filas = list(load_workbook(io.BytesIO(respuesta.content)).active.iter_rows(values_only=True))
        self.assertEqual(len(filas), 2)
        fila = dict(zip(filas[0], filas[1], strict=True))
        self.assertEqual(fila['ID'], self.de_la_programacion.id)
        self.assertEqual(fila['Identificación'], '123')
        self.assertEqual(fila['Empleado'], 'Beatriz')
        self.assertEqual(fila['Nombre concepto'], 'SALARIO')
        self.assertEqual(fila['Días'], 15)
        self.assertEqual(fila['Devengado'], 1000)

    def test_nomina_solo_trae_documentos_de_nomina(self):
        respuesta = self._post(self.lista_documento, informe='nomina')

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        ids = {fila['id'] for fila in respuesta.data['results']}
        self.assertEqual(ids, {self.de_la_programacion.documento_id, self.de_otra.documento_id})

    def test_nomina_excel_filtrado_por_programacion(self):
        from openpyxl import load_workbook

        filtros = [{
            'propiedad': 'programacion_detalle__programacion_id',
            'operador': '=', 'valor': self.programacion.id,
        }]

        respuesta = self._post(self.excel_documento, filtros, informe='nomina')

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('nominas.xlsx', respuesta['Content-Disposition'])
        filas = list(load_workbook(io.BytesIO(respuesta.content)).active.iter_rows(values_only=True))
        self.assertEqual(len(filas), 2)
        fila = dict(zip(filas[0], filas[1], strict=True))
        self.assertEqual(fila['ID'], self.de_la_programacion.documento_id)
        self.assertEqual(fila['Identificación'], '123')
        self.assertEqual(fila['Empleado'], 'Beatriz')
        self.assertEqual(fila['Grupo'], 'Grupo 1')
        self.assertEqual(fila['Aprobado'], 'No')

    def test_lista_de_documentos_filtra_por_programacion(self):
        from general.views.documento import GenDocumentoViewSet

        class VistaLista(GenDocumentoViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        lista = VistaLista.as_view({'post': 'lista'})
        documento = self.de_la_programacion.documento

        for propiedad, valor in (
            ('programacion_detalle__programacion_id', self.programacion.id),
            ('programacion_detalle_id', documento.programacion_detalle_id),
        ):
            filtros = [{'propiedad': propiedad, 'operador': '=', 'valor': valor}]
            respuesta = lista(self.factory.post('/', {'filtros': filtros}, format='json'))

            self.assertEqual(respuesta.status_code, 200, respuesta.data)
            self.assertEqual([fila['id'] for fila in respuesta.data['results']], [documento.id])

    def test_get_de_documento_trae_totales_de_nomina_de_solo_lectura(self):
        from general.models import GenDocumento
        from general.views.documento import GenDocumentoViewSet

        class VistaDocumento(GenDocumentoViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        documento = self.de_la_programacion.documento
        GenDocumento.objects.filter(pk=documento.pk).update(
            devengado=1000, deduccion=80, base_cotizacion=900, base_prestacion=950,
        )

        respuesta = VistaDocumento.as_view({'get': 'retrieve'})(self.factory.get('/'), pk=documento.pk)

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual(
            {campo: Decimal(respuesta.data[campo]) for campo in (
                'devengado', 'deduccion', 'base_cotizacion', 'base_prestacion',
            )},
            {'devengado': 1000, 'deduccion': 80, 'base_cotizacion': 900, 'base_prestacion': 950},
        )

        VistaDocumento.as_view({'patch': 'partial_update'})(
            self.factory.patch('/', {'devengado': 1, 'base_cotizacion': 1}, format='json'), pk=documento.pk,
        )
        documento.refresh_from_db()
        self.assertEqual(documento.devengado, 1000)
        self.assertEqual(documento.base_cotizacion, 900)

    def test_get_de_documento_detalle_trae_concepto_y_valores_de_nomina(self):
        from general.views.documento_detalle import GenDocumentoDetalleViewSet
        from humano.models import HumCredito

        class VistaDetalle(GenDocumentoDetalleViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        type(self.de_la_programacion).objects.filter(pk=self.de_la_programacion.pk).update(
            deduccion=40, base_cotizacion=900, base_prestacion=950, hora=8333.33,
            porcentaje=75,
            credito=HumCredito.objects.create(
                fecha_inicio=date(2026, 1, 1), contrato=self.de_la_programacion.documento.contrato,
            ),
        )

        respuesta = VistaDetalle.as_view({'get': 'retrieve'})(self.factory.get('/'), pk=self.de_la_programacion.pk)

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual(respuesta.data['concepto_id'], 1)
        self.assertEqual(respuesta.data['concepto_nombre'], 'SALARIO')
        self.assertEqual(Decimal(respuesta.data['devengado']), 1000)
        self.assertEqual(Decimal(respuesta.data['deduccion']), 40)
        self.assertEqual(Decimal(respuesta.data['base_cotizacion']), 900)
        self.assertEqual(Decimal(respuesta.data['base_prestacion']), 950)
        self.assertEqual(Decimal(respuesta.data['hora']), Decimal('8333.33'))
        self.assertEqual(Decimal(respuesta.data['porcentaje']), 75)
        self.assertEqual(respuesta.data['credito_id'], respuesta.data['credito'])
        self.assertIsNotNone(respuesta.data['credito_id'])


class ImportarAdicionalTests(TenantTestCase):
    """
    `adicional/importar/`: el front dice si el archivo es de adicionales libres
    (`permanente=true`) o de una programación (`permanente=false` +
    `programacion_id`), y ese modo vale para todas las filas.
    """

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        from humano.models import HumConcepto
        from humano.views.adicional import HumAdicionalViewSet

        class Vista(HumAdicionalViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self.importar = Vista.as_view({'post': 'importar'})
        self.plantilla = Vista.as_view({'get': 'importar_ejemplo'})
        self.factory = APIRequestFactory()

        pais = GenPais.objects.create(id=1, nombre='Colombia')
        ciudad = GenCiudad.objects.create(
            id=1, nombre='Bogotá', estado=GenEstado.objects.create(id=1, nombre='Cundinamarca', pais=pais),
        )
        grupo = HumGrupo.objects.create(nombre='Grupo 1')
        HumPagoTipo.objects.create(id=1, nombre='Nomina')
        self.contrato = HumContrato.objects.create(
            fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 1), grupo=grupo,
            contrato_tipo=HumContratoTipo.objects.create(id=1, nombre='Indefinido'),
            contacto=GenContacto.objects.create(
                numero_identificacion='123', nombre_corto='Beatriz', direccion='x', telefono='1',
                correo='e@e.com', ciudad=ciudad, empleado=True,
                identificacion=GenIdentificacion.objects.create(id=1, nombre='CC'),
                tipo_persona=GenTipoPersona.objects.create(id=1, nombre='Natural'),
            ),
        )
        self.concepto = HumConcepto.objects.create(id=1, nombre='BONIFICACION', adicional=True)
        self.programacion = HumProgramacion.objects.create(
            fecha_desde=date(2026, 3, 1), fecha_hasta=date(2026, 3, 15),
            fecha_hasta_periodo=date(2026, 3, 15), grupo=grupo, pago_tipo_id=1,
        )

    def _archivo(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        from humano.serializers import HumAdicionalImportarSerializer
        from utilidades.mixins import ImportarExcelMixin

        serializer = HumAdicionalImportarSerializer(permanente=True)
        wb = Workbook()
        ws = wb.active
        ws.append([
            ImportarExcelMixin._encabezado_importar(campo, encabezado, serializer.campos_requeridos)
            for campo, encabezado in serializer.campos_excel
        ])
        ws.append([self.contrato.id, self.concepto.id, 50000, 0, 'No', 'Bono'])
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile('adicionales.xlsx', buf.getvalue())

    def _post(self, **datos):
        datos = {'archivo': self._archivo(), **datos}
        return self.importar(self.factory.post('/', datos, format='multipart'))

    def test_importa_libres_como_permanentes_sin_programacion(self):
        from humano.models import HumAdicional

        respuesta = self._post(permanente='true')

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual(respuesta.data, {'creados': 1})
        adicional = HumAdicional.objects.get()
        self.assertTrue(adicional.permanente)
        self.assertIsNone(adicional.programacion_id)
        self.assertEqual(adicional.valor, 50000)

    def test_importa_en_la_programacion(self):
        from humano.models import HumAdicional

        respuesta = self._post(permanente='false', programacion_id=self.programacion.id)

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        adicional = HumAdicional.objects.get()
        self.assertFalse(adicional.permanente)
        self.assertEqual(adicional.programacion_id, self.programacion.id)

    def test_sin_modo_da_400(self):
        from humano.models import HumAdicional

        respuesta = self._post()

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('permanente', respuesta.data['detail'])
        self.assertFalse(HumAdicional.objects.exists())

    def test_no_permanente_sin_programacion_da_400(self):
        respuesta = self._post(permanente='false')

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('programacion_id', respuesta.data)

    def test_permanente_con_programacion_da_400(self):
        respuesta = self._post(permanente='true', programacion_id=self.programacion.id)

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('programacion_id', respuesta.data)

    def test_programacion_inexistente_da_404(self):
        respuesta = self._post(permanente='false', programacion_id=999999)

        self.assertEqual(respuesta.status_code, 404)

    def test_programacion_generada_da_400(self):
        from humano.models import HumAdicional

        HumProgramacion.objects.filter(pk=self.programacion.pk).update(estado_generado=True)

        respuesta = self._post(permanente='false', programacion_id=self.programacion.id)

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('detail', respuesta.data)
        self.assertFalse(HumAdicional.objects.exists())

    def test_la_plantilla_no_trae_la_columna_permanente(self):
        from openpyxl import load_workbook

        respuesta = self.plantilla(self.factory.get('/'))

        self.assertEqual(respuesta.status_code, 200)
        encabezados = next(load_workbook(io.BytesIO(respuesta.content)).active.iter_rows(values_only=True))
        self.assertNotIn('Permanente', encabezados)
