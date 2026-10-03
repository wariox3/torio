from io import BytesIO

from django.http import HttpResponse
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, extend_schema, inline_serializer
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import mixins, serializers, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response

from humano.models import HumProgramacion, HumProgramacionDetalle
from humano.models.programacion_detalle import MENSAJE_PROGRAMACION_CERRADA
from humano.serializers import (
    HumProgramacionDetalleImportarHorasSerializer,
    HumProgramacionDetalleSerializer,
)
from humano.serializers.programacion_detalle_importar_horas import CAMPOS_HORAS
from utilidades.mixins import FiltrosDinamicosMixin, ImportarExcelMixin
from utilidades.mixins.importar_excel import _FUENTE_ENCABEZADO, _FUENTE_NORMAL, _crear_workbook
from utilidades.throttles import ImportarUsuarioTenantThrottle

_PROGRAMACION_PARAM = OpenApiParameter(
    'programacion_id', int, required=True,
    description='Programación cuyos detalles van en la plantilla.',
)

_ImportarHorasRequest = inline_serializer(
    name='ImportarHorasProgramacionRequest',
    fields={
        'archivo': serializers.FileField(help_text='Archivo .xlsx'),
        'programacion_id': serializers.IntegerField(help_text='Programación padre'),
    },
)


class ImportarHorasRequestSerializer(serializers.Serializer):
    programacion_id = serializers.IntegerField()


@extend_schema(tags=['Programación detalle'])
class HumProgramacionDetalleViewSet(
    FiltrosDinamicosMixin,
    ImportarExcelMixin,
    mixins.CreateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.UpdateModelMixin,
    viewsets.GenericViewSet,
):
    serializer_class = HumProgramacionDetalleSerializer

    def get_queryset(self):
        return HumProgramacionDetalle.objects.select_related(
            *HumProgramacionDetalleSerializer.select_related_lista
        )

    # ---- importación de horas ----
    #
    # Los detalles no se crean por Excel (eso lo hace `cargar-contrato`): solo se
    # les cargan las horas. Las dos acciones del mixin se reemplazan por
    # `importar-horas-ejemplo` e `importar-horas`, y la programación la manda el
    # front en `programacion_id`.

    def get_serializer_importar(self):
        return HumProgramacionDetalleImportarHorasSerializer(self.programacion_importar)

    @extend_schema(
        summary='Descargar plantilla de horas',
        description=(
            'Devuelve un .xlsx con un renglón por detalle de la programación: su ID, '
            'la identificación y el nombre del empleado, y las horas actuales. El '
            'usuario edita las horas y sube el mismo archivo a `importar-horas`. No '
            'exige que la programación esté abierta: es una lectura.'
        ),
        parameters=[_PROGRAMACION_PARAM],
        responses={(200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'): OpenApiTypes.BINARY},
    )
    @action(detail=False, methods=['get'], url_path='importar-horas-ejemplo')
    def importar_ejemplo(self, request):
        serializer = ImportarHorasRequestSerializer(data=request.query_params)
        serializer.is_valid(raise_exception=True)

        try:
            programacion = HumProgramacion.objects.get(pk=serializer.validated_data['programacion_id'])
        except HumProgramacion.DoesNotExist:
            raise NotFound('Programación no encontrada.')

        importar = HumProgramacionDetalleImportarHorasSerializer(programacion)
        detalles = (
            HumProgramacionDetalle.objects.filter(programacion=programacion)
            .select_related('contrato__contacto')
            .order_by('contrato__contacto__nombre_corto', 'id')
        )

        wb = _crear_workbook()
        ws = wb.active
        ws.title = 'Datos'
        fondo = PatternFill('solid', fgColor='D9D9D9')
        for col, (campo, encabezado) in enumerate(importar.campos_excel, start=1):
            texto = self._encabezado_importar(campo, encabezado, importar.campos_requeridos)
            celda = ws.cell(row=1, column=col, value=texto)
            celda.font = _FUENTE_ENCABEZADO
            celda.fill = fondo
            ws.column_dimensions[get_column_letter(col)].width = min(max(len(texto), 12) + 2, 50)

        for fila, detalle in enumerate(detalles, start=2):
            contacto = detalle.contrato.contacto
            valores = [
                detalle.id, contacto.numero_identificacion, contacto.nombre_corto,
                *(getattr(detalle, campo) for campo, _ in CAMPOS_HORAS),
            ]
            for col, valor in enumerate(valores, start=1):
                ws.cell(row=fila, column=col, value=valor).font = _FUENTE_NORMAL

        buf = BytesIO()
        wb.save(buf)
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="importar_{importar.nombre_archivo}.xlsx"'
        return response

    @extend_schema(
        summary='Importar horas desde Excel',
        description=(
            'Actualiza las horas de los detalles de la programación indicada en '
            '`programacion_id`. Cada fila trae el ID de un detalle de esa programación; '
            'una celda de horas vacía se guarda como 0. La programación no puede estar '
            'generada ni aprobada. Procesamiento todo-o-nada: si alguna fila falla, no '
            'se guarda nada. La respuesta exitosa es `{creados: N}`, donde N son los '
            'detalles actualizados.'
        ),
        request={'multipart/form-data': _ImportarHorasRequest},
    )
    @action(
        detail=False, methods=['post'], url_path='importar-horas',
        parser_classes=[MultiPartParser],
        throttle_classes=[ImportarUsuarioTenantThrottle],
    )
    def importar(self, request):
        serializer = ImportarHorasRequestSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        try:
            programacion = HumProgramacion.objects.get(pk=serializer.validated_data['programacion_id'])
        except HumProgramacion.DoesNotExist:
            raise NotFound('Programación no encontrada.')

        if programacion.estado_generado or programacion.estado_aprobado:
            return Response({'detail': MENSAJE_PROGRAMACION_CERRADA}, status=status.HTTP_400_BAD_REQUEST)

        self.programacion_importar = programacion
        return super().importar(request)
