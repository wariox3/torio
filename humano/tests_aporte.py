"""
Pruebas del aporte a seguridad social: cargar y eliminar contratos, generar, aprobar,
desaprobar y el plano PILA.

El aporte sale de la nómina aprobada del mes, así que cada prueba arma esa
nómina con el flujo real de la programación (cargar, generar, aprobar) sobre los
catálogos reales: los porcentajes de riesgo, los conceptos y las entidades.
"""
import io
from datetime import date
from decimal import Decimal

from django.core.management import call_command
from django_tenants.test.cases import TenantTestCase
from rest_framework import permissions
from rest_framework.test import APIRequestFactory

from general.models import (
    GenConfiguracion,
    GenContacto,
    GenDocumento,
    GenIdentificacion,
    GenTipoPersona,
)
from humano.models import (
    HumAporte,
    HumAporteContrato,
    HumAporteDetalle,
    HumAporteEntidad,
    HumContrato,
    HumGrupo,
    HumProgramacion,
    HumSucursal,
)
from humano.servicios import (
    AporteError,
    aprobar_aporte,
    aprobar_programacion,
    cargar_contratos,
    cargar_contratos_aporte,
    desaprobar_aporte,
    desgenerar_aporte,
    eliminar_contrato_aporte,
    generar_aporte,
    generar_plano,
    generar_programacion,
)

ENTIDAD_SALUD = 1
ENTIDAD_PENSION = 50
ENTIDAD_CAJA = 70
ENTIDAD_RIESGO = 59
ENTIDAD_SENA = 113
ENTIDAD_ICBF = 112


class FechasAporteTests(TenantTestCase):
    """Las fechas y el periodo de salud salen del año y el mes."""

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def test_mes_de_31(self):
        self.assertEqual(HumAporte.calcular_fechas(2026, 3), {
            'fecha_desde': date(2026, 3, 1), 'fecha_hasta': date(2026, 3, 30),
            'fecha_hasta_periodo': date(2026, 3, 31), 'anio_salud': 2026, 'mes_salud': 4,
        })

    def test_febrero_bisiesto(self):
        fechas = HumAporte.calcular_fechas(2028, 2)
        self.assertEqual((fechas['fecha_hasta'], fechas['fecha_hasta_periodo']), (date(2028, 2, 29),) * 2)

    def test_diciembre_paga_salud_en_enero(self):
        fechas = HumAporte.calcular_fechas(2026, 12)
        self.assertEqual((fechas['anio_salud'], fechas['mes_salud']), (2027, 1))

    def test_la_api_calcula_las_fechas_e_ignora_las_enviadas(self):
        from humano.serializers import HumAporteSerializer

        serializer = HumAporteSerializer(data={
            'anio': 2026, 'mes': 3, 'presentacion': 'S', 'fecha_desde': '2020-01-01', 'mes_salud': 9,
        })
        serializer.is_valid(raise_exception=True)
        aporte = serializer.save()

        self.assertEqual((aporte.fecha_desde, aporte.fecha_hasta), (date(2026, 3, 1), date(2026, 3, 30)))
        self.assertEqual(aporte.mes_salud, 4)

    def test_la_api_rechaza_un_mes_invalido(self):
        from humano.serializers import HumAporteSerializer

        self.assertFalse(HumAporteSerializer(data={'anio': 2026, 'mes': 13, 'presentacion': 'S'}).is_valid())


class AporteTests(TenantTestCase):

    @classmethod
    def setup_tenant(cls, tenant):
        tenant.nombre = 'Test'
        tenant.celular = '0'
        tenant.correo = 'test@test.com'

    def setUp(self):
        from general.signals import limpiar_caches

        self.addCleanup(limpiar_caches)
        call_command('cargar_datos_tenant', schema=self.tenant.schema_name, inicial=True, stdout=io.StringIO())
        GenConfiguracion.objects.filter(pk=1).update(
            hum_factor=Decimal('8'), hum_salario_minimo=Decimal('1423500'),
            hum_auxilio_transporte=Decimal('200000'),
            gen_empresa_nombre_corto='EMPRESA DE PRUEBA', gen_empresa_numero_identificacion='900123456',
            gen_empresa_digito_verificacion='7',
        )
        self.contacto = GenContacto.objects.create(
            numero_identificacion='1036000111', nombre_corto='PEREZ GOMEZ JUAN', nombre1='JUAN',
            apellido1='PEREZ', apellido2='GOMEZ', direccion='Calle 1', telefono='1', correo='e@e.com',
            identificacion=GenIdentificacion.objects.get(aporte='CC'), ciudad_id=1,
            tipo_persona=GenTipoPersona.objects.first(), empleado=True,
        )
        self.grupo = HumGrupo.objects.create(nombre='Grupo aporte')
        self.sucursal = HumSucursal.objects.create(nombre='Principal', codigo='01')
        self.contrato = HumContrato.objects.create(
            fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 1), salario=Decimal('2000000'),
            auxilio_transporte=True, contrato_tipo_id=1, contacto=self.contacto, grupo=self.grupo,
            sucursal=self.sucursal, salud_id=1, pension_id=1, tipo_cotizante_id=1, subtipo_cotizante_id=1,
            riesgo_id=1, entidad_salud_id=ENTIDAD_SALUD, entidad_pension_id=ENTIDAD_PENSION,
            entidad_caja_id=ENTIDAD_CAJA, ciudad_labora_id=1,
        )

    def _nomina_de_marzo(self, aprobar=True):
        programacion = HumProgramacion.objects.create(
            fecha_desde=date(2026, 3, 1), fecha_hasta=date(2026, 3, 30), fecha_hasta_periodo=date(2026, 3, 31),
            dias=30, dias_reales=30, grupo=self.grupo, pago_tipo_id=1,
        )
        cargar_contratos(programacion)
        programacion = generar_programacion(programacion)
        if aprobar:
            aprobar_programacion(programacion)

    def _aporte(self):
        return HumAporte.objects.create(
            anio=2026, mes=3, presentacion='S', sucursal=self.sucursal,
            entidad_riesgo_id=ENTIDAD_RIESGO, entidad_sena_id=ENTIDAD_SENA, entidad_icbf_id=ENTIDAD_ICBF,
            **HumAporte.calcular_fechas(2026, 3),
        )

    def _contactos_entidades(self):
        from humano.models import HumEntidad

        for entidad in HumEntidad.objects.filter(id__in=[ENTIDAD_SALUD, ENTIDAD_PENSION, ENTIDAD_CAJA, ENTIDAD_RIESGO]):
            GenContacto.objects.create(
                numero_identificacion=entidad.numero_identificacion, nombre_corto=entidad.nombre,
                direccion='x', telefono='1', correo='e@e.com', identificacion=GenIdentificacion.objects.get(aporte='CC'),
                ciudad_id=1, tipo_persona=GenTipoPersona.objects.first(),
            )

    # ---- cargar ----

    def test_cargar_toma_el_ibc_de_la_nomina_aprobada(self):
        self._nomina_de_marzo()
        aporte = self._aporte()

        self.assertEqual(cargar_contratos_aporte(aporte), 1)

        aporte_contrato = HumAporteContrato.objects.get(aporte=aporte)
        self.assertEqual(aporte_contrato.dias, 30)
        self.assertEqual(aporte_contrato.base_cotizacion, Decimal('2000000'))
        self.assertEqual(aporte_contrato.cotizacion_pension_empleado, Decimal('80000'))
        self.assertEqual(aporte_contrato.cotizacion_salud_empleado, Decimal('80000'))
        self.assertEqual(aporte_contrato.entidad_riesgo_id, ENTIDAD_RIESGO)
        aporte.refresh_from_db()
        self.assertEqual((aporte.contratos, aporte.empleados), (1, 1))

    def test_la_nomina_sin_aprobar_no_cuenta(self):
        self._nomina_de_marzo(aprobar=False)
        aporte = self._aporte()

        cargar_contratos_aporte(aporte)

        self.assertEqual(HumAporteContrato.objects.get(aporte=aporte).base_cotizacion, 0)

    def test_cargar_dos_veces_no_duplica(self):
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)

        self.assertEqual(cargar_contratos_aporte(aporte), 0)
        self.assertEqual(HumAporteContrato.objects.filter(aporte=aporte).count(), 1)

    # ---- eliminar contrato ----

    def test_eliminar_contrato_recuenta_el_aporte(self):
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)

        eliminar_contrato_aporte(HumAporteContrato.objects.get(aporte=aporte))

        self.assertFalse(HumAporteContrato.objects.filter(aporte=aporte).exists())
        aporte.refresh_from_db()
        self.assertEqual((aporte.contratos, aporte.empleados), (0, 0))

    def test_no_elimina_contrato_con_el_aporte_generado(self):
        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        generar_aporte(aporte)

        with self.assertRaises(AporteError):
            eliminar_contrato_aporte(HumAporteContrato.objects.get(aporte=aporte))
        self.assertTrue(HumAporteContrato.objects.filter(aporte=aporte).exists())

    def test_endpoint_eliminar_contrato(self):
        from humano.views.aporte_contrato import HumAporteContratoViewSet

        class Vista(HumAporteContratoViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        aporte_contrato = HumAporteContrato.objects.get(aporte=aporte)
        vista = Vista.as_view({'delete': 'destroy'})

        generar_aporte(aporte)
        respuesta = vista(APIRequestFactory().delete('/'), pk=aporte_contrato.pk)
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('detail', respuesta.data)

        desgenerar_aporte(aporte)
        respuesta = vista(APIRequestFactory().delete('/'), pk=aporte_contrato.pk)
        self.assertEqual(respuesta.status_code, 204)
        self.assertFalse(HumAporteContrato.objects.filter(pk=aporte_contrato.pk).exists())

    # ---- generar ----

    def test_generar(self):
        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)

        aporte = generar_aporte(aporte)

        detalle = HumAporteDetalle.objects.get(aporte_contrato__aporte=aporte)
        self.assertEqual(detalle.dias_pension, 30)
        self.assertEqual(detalle.cotizacion_pension, Decimal('320000'))   # 16 %
        self.assertEqual(detalle.cotizacion_salud, Decimal('80000'))      # 4 %
        self.assertEqual(detalle.cotizacion_riesgos, Decimal('10500'))    # 0,522 % → 10.440, a la centena
        self.assertEqual(detalle.cotizacion_caja, Decimal('80000'))       # 4 %
        self.assertEqual(aporte.cotizacion_total, Decimal('490500'))
        self.assertEqual(aporte.base_cotizacion, Decimal('2000000'))
        # La empresa pone lo que no se le descontó al empleado.
        self.assertEqual(aporte.cotizacion_pension_empresa, Decimal('240000'))
        self.assertEqual(aporte.cotizacion_salud_empresa, Decimal('0'))
        self.assertTrue(aporte.estado_generado)

        entidades = dict(HumAporteEntidad.objects.filter(aporte=aporte).values_list('tipo', 'cotizacion'))
        self.assertEqual(entidades, {
            'PENSION': Decimal('320000'), 'SALUD': Decimal('80000'), 'CAJA': Decimal('80000'),
            'RIESGO': Decimal('10500'), 'SENA': Decimal('0'), 'ICBF': Decimal('0'),
        })

    def test_sin_contratos_no_genera(self):
        with self.assertRaises(AporteError):
            generar_aporte(self._aporte())

    def test_desgenerar_deja_todo_en_cero(self):
        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)

        aporte = desgenerar_aporte(generar_aporte(aporte))

        self.assertFalse(aporte.estado_generado)
        self.assertEqual(aporte.cotizacion_total, 0)
        self.assertFalse(HumAporteDetalle.objects.exists())
        self.assertFalse(HumAporteEntidad.objects.exists())
        self.assertEqual(HumAporteContrato.objects.get(aporte=aporte).cotizacion_total, 0)

    # ---- aprobar y desaprobar ----

    def test_no_aprueba_si_las_entidades_no_son_contacto(self):
        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        aporte = generar_aporte(aporte)

        with self.assertRaises(AporteError) as error:
            aprobar_aporte(aporte)

        self.assertEqual({e['tipo'] for e in error.exception.errores}, {'PENSION', 'SALUD', 'CAJA', 'RIESGO'})
        self.assertFalse(GenDocumento.objects.filter(aporte=aporte).exists())

    def test_aprobar_crea_un_documento_por_entidad_y_desaprobar_los_borra(self):
        self._nomina_de_marzo()
        self._contactos_entidades()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)

        aporte = aprobar_aporte(generar_aporte(aporte))

        documentos = {d.orden_compra: d for d in GenDocumento.objects.filter(aporte=aporte)}
        self.assertEqual(set(documentos), {'PENSION', 'SALUD', 'CAJA', 'RIESGO'})
        pension = documentos['PENSION']
        self.assertEqual(pension.documento_tipo_id, 22)
        self.assertTrue(pension.estado_aprobado)
        self.assertIsNotNone(pension.numero)
        self.assertEqual((pension.total, pension.subtotal), (Decimal('320000'), Decimal('240000')))
        self.assertEqual(pension.pendiente, pension.total)
        self.assertEqual(len({d.numero for d in documentos.values()}), 4)

        aporte = desaprobar_aporte(aporte)

        self.assertFalse(aporte.estado_aprobado)
        self.assertFalse(GenDocumento.objects.filter(aporte=aporte).exists())

    def test_no_desaprueba_si_ya_tiene_egreso(self):
        self._nomina_de_marzo()
        self._contactos_entidades()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        aporte = aprobar_aporte(generar_aporte(aporte))
        GenDocumento.objects.filter(aporte=aporte, orden_compra='SALUD').update(afectado=Decimal('80000'))

        with self.assertRaises(AporteError):
            desaprobar_aporte(aporte)
        self.assertEqual(GenDocumento.objects.filter(aporte=aporte).count(), 4)

    # ---- plano ----

    def test_plano_pila(self):
        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        aporte = generar_aporte(aporte)

        contenido, nombre = generar_plano(aporte)

        encabezado, cotizante = contenido.decode('cp1252').splitlines()
        self.assertTrue(nombre.startswith('pila') and nombre.endswith('.txt'))
        # Anchos fijos de la Resolución 2388: 359 el encabezado y 693 cada cotizante.
        self.assertEqual(len(encabezado), 359)
        self.assertEqual(len(cotizante), 693)
        self.assertEqual(encabezado[209:225].strip(), '900123456')
        self.assertEqual(encabezado[304:318], '2026-032026-04')
        self.assertEqual(cotizante[:9], '0200001CC')
        self.assertEqual(cotizante[9:25].strip(), '1036000111')
        self.assertEqual(cotizante[201:210], '002000000')  # IBC pensión
        self.assertEqual(cotizante[237:244], '0.16000')    # Tarifa pensión
        self.assertEqual(cotizante[244:253], '000320000')  # Cotización pensión

    def test_plano_sin_generar(self):
        from humano.servicios import PilaError

        with self.assertRaises(PilaError):
            generar_plano(self._aporte())

    def test_imprimir(self):
        from humano.views.aporte import HumAporteViewSet

        class Vista(HumAporteViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        generar_aporte(aporte)
        vista = Vista.as_view({'post': 'imprimir'}, **Vista.imprimir.kwargs)

        respuesta = vista(APIRequestFactory().post('/', {'aporte_id': aporte.id}, format='json'))

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertTrue(respuesta.content.startswith(b'%PDF'))

    def test_excel_de_contratos_detalles_y_entidades(self):
        from openpyxl import load_workbook

        from humano.views.aporte_contrato import HumAporteContratoViewSet
        from humano.views.aporte_detalle import HumAporteDetalleViewSet
        from humano.views.aporte_entidad import HumAporteEntidadViewSet

        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        generar_aporte(aporte)

        casos = (
            (HumAporteContratoViewSet, 'aporte_id', 'Empleado', 1),
            (HumAporteDetalleViewSet, 'aporte_contrato__aporte_id', 'Empleado', HumAporteDetalle.objects.count()),
            (HumAporteEntidadViewSet, 'aporte_id', 'Nombre', HumAporteEntidad.objects.count()),
        )
        for clase, propiedad, columna, filas in casos:
            with self.subTest(clase.__name__):
                vista = type('Vista', (clase,), {
                    'authentication_classes': [], 'permission_classes': [permissions.AllowAny],
                    'throttle_classes': [],
                }).as_view({'post': 'excel'})
                filtro = {'propiedad': propiedad, 'operador': '=', 'valor': aporte.id}

                respuesta = vista(APIRequestFactory().post('/', {'filtros': [filtro]}, format='json'))

                self.assertEqual(respuesta.status_code, 200)
                hoja = load_workbook(io.BytesIO(respuesta.content)).active
                self.assertIn(columna, [celda.value for celda in hoja[1]])
                self.assertEqual(hoja.max_row - 1, filas)

    def test_endpoint_plano_operador(self):
        from humano.views.aporte import HumAporteViewSet

        class Vista(HumAporteViewSet):
            authentication_classes = []
            permission_classes = [permissions.AllowAny]
            throttle_classes = []

        self._nomina_de_marzo()
        aporte = self._aporte()
        cargar_contratos_aporte(aporte)
        generar_aporte(aporte)
        vista = Vista.as_view({'post': 'plano_operador'}, **Vista.plano_operador.kwargs)

        respuesta = vista(APIRequestFactory().post('/', {'aporte_id': aporte.id}, format='json'))

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'text/plain; charset=windows-1252')
        self.assertIn('attachment; filename="pila', respuesta['Content-Disposition'])
